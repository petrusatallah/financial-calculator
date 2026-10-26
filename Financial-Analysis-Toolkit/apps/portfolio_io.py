"""Portfolio-only Excel and CSV input adapters.

Every imported numeric unit is explicit. Summary-only workbooks contain estimates;
Investments exports with a Returns Input sheet contain dated observations.
"""
from __future__ import annotations

import csv
import io
import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from financial_toolkit.validation import InputError


def _date(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    try:
        return date.fromisoformat(str(value).strip()).isoformat()
    except ValueError as exc:
        raise InputError("Dates must use YYYY-MM-DD.") from exc


def _number(value, label):
    if isinstance(value, bool) or value is None:
        raise InputError(f"{label} must be a number.")
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value:
            raise InputError(f"{label} is missing.")
        if value.endswith("%"):
            value = value[:-1].strip()
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise InputError(f"{label} must be a number.") from exc
    if not math.isfinite(number):
        raise InputError(f"{label} must be finite.")
    return number


def _excel_percent_format(number_format):
    """Excel stores a percentage-formatted numeric cell as a decimal."""
    unquoted = re.sub(r'"(?:[^"]|"")*"|\\.', '', number_format or '')
    return '%' in unquoted


def _summary(ws):
    if ws["A1"].value != "Investment Analysis" or ws["A5"].value != "Return frequency" or ws["F4"].value != "Portfolio weights":
        raise InputError("This is not a compatible Investments Summary workbook.")
    frequency = str(ws["B5"].value).strip().lower()
    if frequency not in ("monthly", "yearly"):
        raise InputError("The Investments Summary has an unsupported frequency.")
    weights = {}
    for row in range(5, ws.max_row+1):
        name = ws.cell(row, 6).value
        if isinstance(name, str) and name.startswith("Total"):
            break
        if name is None:
            raise InputError("The Investments Summary weight table is incomplete.")
        name = str(name).strip()
        if not name or name in weights:
            raise InputError("Investment names in the Summary must be unique.")
        weights[name] = _number(ws.cell(row, 8).value, f"Weight for {name}")
    names = list(weights)
    if len(names) < 2:
        raise InputError("The Summary needs at least two investments.")
    if not math.isclose(sum(weights.values()), 1, abs_tol=1e-7):
        raise InputError("The Summary weights must total 100% to import as current weights.")
    stats_header = next((r for r in range(1, ws.max_row+1)
                         if ws.cell(r, 1).value == "Investment" and str(ws.cell(r, 2).value).startswith("Average (")), None)
    if stats_header is None:
        raise InputError("The Summary's individual statistics table is missing.")
    means = {}
    for index, name in enumerate(names, stats_header+1):
        if ws.cell(index, 1).value != name:
            raise InputError("Summary statistics names do not match portfolio weights.")
        means[name] = _number(ws.cell(index, 2).value, f"Average for {name}")
    matrix_header = next((r for r in range(1, ws.max_row+1)
                          if isinstance(ws.cell(r, 1).value, str)
                          and ws.cell(r, 1).value.startswith("Covariance matrix (decimal return squared)")), None)
    if matrix_header is None:
        raise InputError("The Summary covariance matrix is missing or its units are unclear.")
    if [ws.cell(matrix_header+1, col).value for col in range(2, 2+len(names))] != names:
        raise InputError("Covariance columns do not match investment names.")
    covariance = []
    for i, name in enumerate(names):
        row = matrix_header+2+i
        if ws.cell(row, 1).value != name:
            raise InputError("Covariance rows do not match investment names.")
        covariance.append([_number(ws.cell(row, col).value, f"Covariance {name}")
                           for col in range(2, 2+len(names))])
    return {"kind": "estimates", "source": "Investments Summary", "frequency": frequency,
            "names": names, "expected_returns": means, "covariance": covariance,
            "current_weights": weights, "dates": [], "returns": [],
            "note": "This Summary contains arithmetic averages and covariance, not raw return history."}


def parse_portfolio_file(data: bytes, filename: str, *, scale: str, purpose: str):
    """Parse a compatible Summary or a declared dated-return table.

    `scale` is 'auto', 'percent' (10 means 10%), or 'decimal' (.10 means 10%).
    Explicit percent strings and Excel percentage-formatted cells take precedence.
    `purpose` is 'investments' or 'portfolios'; UI confirmation identifies columns.
    """
    if not isinstance(data, bytes) or not data or len(data) > 20_000_000:
        raise InputError("Choose a nonempty file under 20 MB.")
    if scale not in ("auto", "percent", "decimal") or purpose not in ("investments", "portfolios"):
        raise InputError("Declare input scale and column meaning.")
    ext = Path(filename).suffix.lower()
    summary_data = None
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
            records = list(csv.reader(io.StringIO(text)))
            percent_formats = None
        except (UnicodeDecodeError, csv.Error) as exc:
            raise InputError("CSV must be UTF-8 with a header row.") from exc
    elif ext == ".xlsx":
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise InputError("The Excel workbook could not be read.") from exc
        try:
            investments_summary = ("Summary" in book.sheetnames and
                                   book["Summary"]["A1"].value == "Investment Analysis")
            if purpose == "investments" and investments_summary:
                summary_data = _summary(book["Summary"])
                if book.sheetnames == ["Summary"]:
                    return summary_data
                if len(book.sheetnames) != 2 or "Returns Input" not in book.sheetnames:
                    raise InputError("The Investments export must contain only Summary and Returns Input.")
                sheet = book["Returns Input"]
                summary_periods = _number(book["Summary"]["B8"].value, "Number of periods")
            elif investments_summary:
                raise InputError("Investments Summary data is individual investments, not complete portfolio return histories.")
            else:
                sheet = book.active
            cell_rows = list(sheet.iter_rows(values_only=False))
            records = [[cell.value for cell in row] for row in cell_rows]
            percent_formats = [[_excel_percent_format(cell.number_format) for cell in row]
                               for row in cell_rows]
        finally:
            book.close()
    else:
        raise InputError("Supported formats are .xlsx and .csv.")
    if len(records) < 3:
        raise InputError("A dated return table needs a header and at least two observations.")
    header = [str(value).strip() if value is not None else "" for value in records[0]]
    if header[0].lower() != "date" or len(header) < 3 or any(not name for name in header[1:]) or len(set(header[1:])) != len(header)-1:
        raise InputError("The first column must be Date, followed by unique investment or portfolio names.")
    names = header[1:]
    dates, values, original = [], [], []
    for row_number, row in enumerate(records[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if len(row) > len(header) and any(value is not None and str(value).strip() for value in row[len(header):]):
            raise InputError(f"Row {row_number} has unexpected extra data.")
        if len(row) < len(header):
            raise InputError(f"Row {row_number} has a missing return.")
        dates.append(_date(row[0]))
        raw = list(row[1:len(header)])
        original.append(raw)
        numbers = []
        for column, (name, cell) in enumerate(zip(names, raw), 1):
            explicit_percent = isinstance(cell, str) and cell.strip().endswith("%")
            formatted_percent = (percent_formats is not None
                                 and percent_formats[row_number - 1][column]
                                 and not isinstance(cell, str))
            number = _number(cell, f"{name} on row {row_number}")
            if explicit_percent:
                decimal = number / 100
            elif formatted_percent:
                decimal = number
            elif scale == "auto":
                raise InputError(
                    f"{name} on row {row_number} has no percent sign or Excel percent format. "
                    "Choose Percent or Decimal for unmarked values; their units cannot be inferred safely."
                )
            else:
                decimal = number / 100 if scale == "percent" else number
            if decimal < -1:
                raise InputError(f"{name} on row {row_number} is below -100%.")
            numbers.append(decimal)
        values.append(numbers)
    if len(dates) < 2 or dates != sorted(set(dates)):
        raise InputError("At least two unique, increasing dates are required.")
    parsed = {"kind": "history", "source": filename, "frequency": None, "names": names,
              "dates": dates, "returns": values, "original_values": original,
              "original_scale": scale, "purpose": purpose,
              "normalized_scale": "decimal",
              "note": "Columns are treated as complete portfolios." if purpose == "portfolios" else
                      "Columns are treated as individual investments."}
    if summary_data is not None:
        if names != summary_data["names"] or len(dates) != summary_periods:
            raise InputError("Returns Input must match the Summary investment names and period count.")
        parsed["frequency"] = summary_data["frequency"]
        parsed["current_weights"] = summary_data["current_weights"]
        parsed["note"] = "Dated investment returns from Returns Input; current weights from Summary."
    return parsed
