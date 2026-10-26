"""Excel exports for the two Portfolio workflows only."""
from __future__ import annotations

import io
from datetime import date
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "103C70"
BLUE = "1456A4"
PANEL = "EAF2FA"
BACKGROUND = "FFFFFF"
INK = "172B42"
MUTED = "4A627C"
GRID = "C6D6E6"
GREEN = "B9DDBC"
ORANGE = "FFD09A"
PALE_BLUE = "C9E0F3"
PERCENT = "0.0000%;[Red]-0.0000%;0.0000%"
RATIO = "0.0000;[Red]-0.0000;0.0000"
VARIANCE = "0.00000000;[Red]-0.00000000;0.00000000"
TINY = "0.00000000E+00;[Red]-0.00000000E+00;0.00000000E+00"


def _variance_format(value):
    return TINY if value != 0 and abs(value) < 5e-9 else VARIANCE


def _header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[row].height = 30


def _sheet(book, name, title):
    ws = book.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    for row in (1, 2):
        for col in range(1, 11):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:J2")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 16
    ws.column_dimensions["A"].width = 26
    for letter in "BCDEFGHIJKLMNOP":
        ws.column_dimensions[letter].width = 19
    return ws


def _report_cell(ws, row, col, value, number_format=None, *, bold=False):
    cell = ws.cell(row, col)
    if value is None:
        cell.value = "Unavailable"
        cell.data_type = "s"
    elif isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    if number_format and value is not None and not isinstance(value, str):
        cell.number_format = number_format
    cell.font = Font(name="Aptos", size=10, bold=bold, color=INK)
    cell.fill = PatternFill("solid", fgColor=BACKGROUND if row % 2 else PANEL)
    cell.alignment = Alignment(vertical="center", horizontal="center" if value is None else
                               "right" if isinstance(value, (int, float)) else "left",
                               indent=1 if isinstance(value, str) else 0)
    cell.border = Border(bottom=Side(style="hair", color=GRID))
    return cell


def _report_section(ws, row, label, color=PALE_BLUE):
    for col in range(1, 13):
        ws.cell(row, col).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    cell = _report_cell(ws, row, 1, label, bold=True)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name="Aptos", size=11, bold=True, color=INK)
    ws.row_dimensions[row].height = 23


def _report_table(ws, row, col, headers, rows, formats=None):
    formats = formats or {}
    for offset, label in enumerate(headers):
        cell = _report_cell(ws, row, col + offset, label, bold=True)
        cell.fill = PatternFill("solid", fgColor=PANEL)
        cell.font = Font(name="Aptos", size=9, bold=True, color=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    for index, values in enumerate(rows, row + 1):
        for offset, value in enumerate(values):
            _report_cell(ws, index, col + offset, value, formats.get(offset), bold=False)
        ws.row_dimensions[index].height = 19
    return row + len(rows)


def _report_matrix(ws, row, col, title, names, values, *, correlation=False):
    _report_cell(ws, row, col, title, bold=True).font = Font(name="Aptos", size=10, bold=True, color=MUTED)
    end = _report_table(ws, row + 1, col, ["", *names],
                        [[name, *values[i]] for i, name in enumerate(names)],
                        {j: RATIO if correlation else VARIANCE for j in range(1, len(names) + 1)})
    if not correlation:
        for i, matrix_row in enumerate(values):
            for j, value in enumerate(matrix_row):
                ws.cell(row + 2 + i, col + 1 + j).number_format = _variance_format(value)
    else:
        first = f"{get_column_letter(col + 1)}{row + 2}"
        last = f"{get_column_letter(col + len(names))}{row + 1 + len(names)}"
        ws.conditional_formatting.add(f"{first}:{last}", ColorScaleRule(
            start_type="num", start_value=-1, start_color="E66A6A",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=1, end_color="4389E4"))
    return end


def _report_metric(ws, row, start, end, label, value, fmt, color):
    for at in (row, row + 1):
        for col in range(start, end + 1):
            ws.cell(at, col).fill = PatternFill("solid", fgColor=color)
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    ws.merge_cells(start_row=row + 1, start_column=start, end_row=row + 1, end_column=end)
    label_cell = _report_cell(ws, row, start, label)
    label_cell.fill = PatternFill("solid", fgColor=color)
    label_cell.font = Font(name="Aptos", size=9, bold=True, color=INK)
    value_cell = _report_cell(ws, row + 1, start, value, fmt, bold=True)
    value_cell.fill = PatternFill("solid", fgColor=color)
    value_cell.font = Font(name="Aptos", size=16, bold=True, color=INK)
    ws.row_dimensions[row + 1].height = 29


def _report_base(book, kind, frequency):
    ws = book.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 25
    for letter in "BCDEFGHIJKL":
        ws.column_dimensions[letter].width = 16
    for row in (1, 2):
        for col in range(1, 13):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:L2")
    title = _report_cell(ws, 1, 1, kind + " Report")
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Aptos Display", size=21, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 17
    _report_cell(ws, 3, 1, "Report date", bold=True)
    _report_cell(ws, 3, 2, date.today(), "yyyy-mm-dd")
    _report_cell(ws, 3, 6, "Return frequency", bold=True)
    _report_cell(ws, 3, 7, frequency.title())
    return ws


def _finish_report(ws, last_row):
    fill = PatternFill("solid", fgColor=BACKGROUND)
    for row in ws.iter_rows(min_row=3, max_row=last_row, min_col=1, max_col=12):
        for cell in row:
            if cell.fill.patternType is None:
                cell.fill = fill
    ws.print_area = f"A1:L{last_row}"


def _style_details(book, dashboard):
    dark = PatternFill("solid", fgColor=BACKGROUND)
    panel = PatternFill("solid", fgColor=PANEL)
    for ws in book.worksheets:
        if ws is dashboard:
            continue
        ws.sheet_properties.tabColor = "35516F"
        last_row = max(ws.max_row, 18 if ws._charts else 1)
        last_col = max(ws.max_column, 10 if ws._charts else 1)
        for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=last_col):
            for cell in row:
                if cell.fill.patternType is None:
                    cell.fill = panel if cell.value is not None and row[0].row % 2 == 0 else dark
                if cell.value is not None and cell.coordinate != "A1" and cell.font.color is None:
                    cell.font = Font(name="Aptos", size=10, color=INK)
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws["A1"].font = Font(name="Aptos", size=18, bold=True, color="FFFFFF")


def _markowitz_dashboard(book, inputs, analysis, solution, frontier, baseline):
    names = inputs["names"]
    d, optimum = analysis.diagnostics, solution.diagnostics
    ws = _report_base(book, "Markowitz Optimization", inputs["frequency"])
    _report_section(ws, 5, "1. Investment data summary", GREEN)
    summary_rows = [[name, d["expected_asset_returns"][name],
                     d["current_weights"][name] if d["current_weights"] else None,
                     d["asset_volatility"][name]] for name in names]
    summary_end = _report_table(ws, 7, 1,
                                ["Investment", "Expected return", "Current weight", "Volatility"],
                                summary_rows, {1: PERCENT, 2: PERCENT, 3: PERCENT})
    row = summary_end + 2
    _report_section(ws, row, "2. Optimization settings and constraints", PALE_BLUE)
    settings = [
        ["Objective", inputs.get("objective", "Optimization")],
        ["Input basis", "Historical returns" if inputs["kind"] == "history" else "Expected returns and covariance"],
        ["Risk-free return per period", d.get("risk_free_period_return")],
        ["Source", inputs.get("source", "Manual entry")],
    ]
    constraints = inputs.get("constraints") or {}
    if constraints.get("target_return") is not None:
        settings.append(["Target return floor per period", constraints["target_return"]])
    for label, key in (("Allow short selling", "allow_short"), ("Allow cash or funding", "allow_cash"),
                       ("Gross exposure cap", "gross_cap"), ("Minimum net risky exposure", "net_min"),
                       ("Maximum net risky exposure", "net_max"),
                       ("Annual cash yield", "cash_annual_yield"),
                       ("Annual funding rate", "funding_annual_rate")):
        if key in constraints:
            value = constraints[key]
            if isinstance(value, bool):
                value = "Yes" if value else "No"
            settings.append([label, value])
    settings.append(["Custom groups", len(constraints.get("group_constraints", []))])
    settings_end = _report_table(ws, row + 2, 1, ["Setting", "Applied value"], settings)
    for setting_row, (label, _) in enumerate(settings, row + 3):
        if label in ("Risk-free return per period", "Target return floor per period",
                     "Gross exposure cap", "Minimum net risky exposure",
                     "Maximum net risky exposure", "Annual cash yield", "Annual funding rate"):
            ws.cell(setting_row, 2).number_format = PERCENT
    row = settings_end + 2
    lower, upper = constraints.get("lower_bounds"), constraints.get("upper_bounds")
    if lower is not None and upper is not None:
        _report_cell(ws, row, 1, "Individual investment limits (%)", bold=True)
        row = _report_table(ws, row + 1, 1,
                            ["Investment", "Minimum weight", "Maximum weight"],
                            [[name, lower[name], upper[name]] for name in names],
                            {1: PERCENT, 2: PERCENT}) + 2
    groups = constraints.get("group_constraints") or []
    if groups:
        _report_cell(ws, row, 1, "Custom group limits (%)", bold=True)
        row = _report_table(ws, row + 1, 1,
                            ["Group", "Investments", "Minimum weight", "Maximum weight"],
                            [[group["name"], ", ".join(group["investments"]),
                              group.get("minimum") if group.get("minimum") is not None else "No minimum",
                              group["maximum"]] for group in groups],
                            {2: PERCENT, 3: PERCENT}) + 2
    _report_section(ws, row, "3. Optimized portfolio weights", GREEN)
    weights = [[name,
                d["current_weights"][name] if d["current_weights"] else None,
                optimum["weights"][name],
                (optimum["weights"][name] - d["current_weights"][name]) * 100 if d["current_weights"] else None,
                baseline.diagnostics["weights"][name] if baseline else None] for name in names]
    weights_end = _report_table(ws, row + 2, 1,
                                ["Investment", "Current weight", "Optimized weight", "Change (percentage points)", "Baseline weight"],
                                weights, {1: PERCENT, 2: PERCENT, 3: "+0.0000;-0.0000;0.0000", 4: PERCENT})
    row = weights_end + 2
    _report_section(ws, row, "4. Portfolio statistics", ORANGE)
    _report_metric(ws, row + 2, 1, 3, "Expected return", optimum["expected_return"], PERCENT, "DDF0DF")
    _report_metric(ws, row + 2, 4, 6, "Variance", optimum["variance"],
                   _variance_format(optimum["variance"]), "DBEAF8")
    _report_metric(ws, row + 2, 7, 9, "Volatility", optimum["volatility"], PERCENT, "E5DDF7")
    _report_metric(ws, row + 2, 10, 12, "Sharpe ratio", optimum["estimated_sharpe"], RATIO, "FFF0CA")
    row += 6
    _report_section(ws, row, "5. Covariance and correlation", PALE_BLUE)
    covariance_end = _report_matrix(ws, row + 2, 1, "Covariance (decimal return squared)",
                                    names, d["covariance"])
    if len(names) <= 4:
        correlation_end = _report_matrix(ws, row + 2, 7, "Correlation", names, d["correlation"],
                                         correlation=True)
    else:
        correlation_end = _report_matrix(ws, covariance_end + 2, 1, "Correlation", names,
                                         d["correlation"], correlation=True)
    row = max(covariance_end, correlation_end) + 2
    _report_section(ws, row, "6. Efficient frontier", PALE_BLUE)
    if frontier is not None and frontier.ok and len(frontier.diagnostics["points"]) > 1:
        chart_ws = book["Frontier"]
        chart = ScatterChart()
        chart.style = 10
        chart.title = "Efficient frontier"
        chart.x_axis.title = "Volatility (%)"
        chart.y_axis.title = "Expected return (%)"
        chart.x_axis.numFmt = "0.0%"
        chart.y_axis.numFmt = "0.0%"
        last = 3 + len(frontier.diagnostics["points"])
        chart.series.append(Series(Reference(chart_ws, min_col=2, min_row=4, max_row=last),
                                   Reference(chart_ws, min_col=1, min_row=4, max_row=last),
                                   title="Efficient frontier"))
        chart.series[0].graphicalProperties.line.solidFill = BLUE
        chart.width, chart.height = 22, 10
        ws.add_chart(chart, f"A{row + 2}")
        last_row = row + 19
    else:
        _report_cell(ws, row + 2, 1, "No distinct frontier chart is available for these inputs.")
        last_row = row + 4
    warnings = tuple(dict.fromkeys((*analysis.warnings, *solution.warnings,
                                    *(frontier.warnings if frontier is not None else ()))))
    if warnings:
        _report_cell(ws, last_row + 1, 1, "Estimation warnings", bold=True)
        for offset, warning in enumerate(warnings, last_row + 2):
            ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=12)
            _report_cell(ws, offset, 1, warning)
        last_row += 2 + len(warnings)
    _finish_report(ws, last_row)
    return ws


def _compare_dashboard(book, inputs, analysis):
    names = inputs["names"]
    d = analysis.diagnostics
    ws = _report_base(book, "Compare Portfolios", inputs["frequency"])
    _report_cell(ws, 4, 1, "Risk-free return per period", bold=True)
    _report_cell(ws, 4, 2, d.get("risk_free_period_return"), PERCENT)
    _report_section(ws, 5, "1. Compared portfolios", GREEN)
    source = inputs.get("source", "Manual entry")
    definitions_end = _report_table(ws, 7, 1, ["Portfolio", "Source", "Periods"],
                                    [[name, source, len(inputs["dates"])] for name in names],
                                    {2: "0"})
    row = definitions_end + 2
    _report_section(ws, row, "2. Performance comparison", ORANGE)
    performance_end = _report_table(
        ws, row + 2, 1,
        ["Portfolio", "Average return", "Cumulative return", "Sample variance", "Volatility", "Sharpe ratio"],
        [[item["name"], item["mean"], item["cumulative"], item["variance"],
          item["volatility"], item["sharpe"]] for item in d["comparison"]],
        {1: PERCENT, 2: PERCENT, 3: VARIANCE, 4: PERCENT, 5: RATIO})
    for result_row, item in enumerate(d["comparison"], row + 3):
        ws.cell(result_row, 4).number_format = _variance_format(item["variance"])
    row = performance_end + 2
    _report_section(ws, row, "3. Portfolio covariance and correlation", PALE_BLUE)
    covariance_end = _report_matrix(ws, row + 2, 1, "Covariance (decimal return squared)",
                                    names, d["covariance"])
    if len(names) <= 4:
        correlation_end = _report_matrix(ws, row + 2, 7, "Correlation", names, d["correlation"],
                                         correlation=True)
    else:
        correlation_end = _report_matrix(ws, covariance_end + 2, 1, "Correlation", names,
                                         d["correlation"], correlation=True)
    row = max(covariance_end, correlation_end) + 2
    _report_section(ws, row, "4. Cumulative returns comparison", PALE_BLUE)
    if len(inputs["dates"]) > 1:
        chart_ws = book["Performance"]
        chart = LineChart()
        chart.style = 10
        chart.title = "Cumulative return"
        chart.y_axis.title = "Cumulative return (%)"
        chart.add_data(Reference(chart_ws, min_col=2, max_col=len(names) + 1,
                                 min_row=3, max_row=3 + len(inputs["dates"])), titles_from_data=True)
        for series, name in zip(chart.series, names):
            series.title = SeriesLabel(v=name)
        for index, series in enumerate(chart.series):
            series.graphicalProperties.line.solidFill = (BLUE, "168B50", "E68A2E", "8A63B8")[index % 4]
        chart.set_categories(Reference(chart_ws, min_col=1, min_row=4,
                                       max_row=3 + len(inputs["dates"])))
        chart.y_axis.numFmt = "0%"
        chart.width, chart.height = 22, 10
        ws.add_chart(chart, f"A{row + 2}")
        last_row = row + 19
    else:
        _report_cell(ws, row + 2, 1, "The chart requires at least two periods.")
        last_row = row + 4
    for offset, warning in enumerate(analysis.warnings, last_row + 1):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=12)
        _report_cell(ws, offset, 1, warning)
    last_row += len(analysis.warnings)
    _finish_report(ws, last_row)
    return ws


def portfolio_excel_bytes(kind, inputs, analysis, solution=None, frontier=None, baseline=None):
    book = Workbook()
    book.remove(book.active)
    names = inputs["names"]
    source = _sheet(book, "Inputs", f"{kind} — validated inputs")
    source["A3"], source["B3"] = "Source", inputs.get("source", "Manual entry")
    source["A4"], source["B4"] = "Frequency", inputs["frequency"].title()
    source["A5"], source["B5"] = "Input type", inputs["kind"]
    source["A6"], source["B6"] = "Return scale", "Decimal internally; displayed as percent"
    if analysis.diagnostics.get("risk_free_period_return") is not None:
        source["A7"], source["B7"] = "Risk-free rate per period", analysis.diagnostics["risk_free_period_return"]
        source["B7"].number_format = PERCENT
    source["A8"] = "Objective / basis"
    source["B8"] = inputs.get("objective", "Historical comparison")
    if inputs["kind"] == "history":
        _header(source, 9, ["Date", *names])
        for row_no, (day, values) in enumerate(zip(inputs["dates"], inputs["returns"]), 10):
            source.cell(row_no, 1, day)
            for col, value in enumerate(values, 2):
                source.cell(row_no, col, value).number_format = PERCENT
    else:
        _header(source, 9, ["Investment", "Expected return per period", "Current weight"])
        for row_no, name in enumerate(names, 10):
            source.cell(row_no, 1, name)
            source.cell(row_no, 2, analysis.diagnostics["expected_asset_returns"][name]).number_format = PERCENT
            if analysis.diagnostics["current_weights"] is not None:
                source.cell(row_no, 3, analysis.diagnostics["current_weights"][name]).number_format = PERCENT
    if inputs.get("constraints"):
        row_no = max(13, 11+len(inputs.get("dates", [])), 11+len(names))
        _header(source, row_no, ["Applied constraint", "Value (decimal)"])
        for label, value in inputs["constraints"].items():
            row_no += 1
            source.cell(row_no, 1, label)
            source.cell(row_no, 2, str(value))
    results = _sheet(book, "Results", f"{kind} — results")
    matrix = _sheet(book, "Matrices", "Sample or supplied covariance and correlation")
    d = analysis.diagnostics
    _header(matrix, 3, ["Covariance (decimal return squared)", *names])
    for i, name in enumerate(names):
        matrix.cell(i+4, 1, name)
        for j, value in enumerate(d["covariance"][i], 2):
            matrix.cell(i+4, j, value).number_format = _variance_format(value)
    offset = len(names)+6
    _header(matrix, offset, ["Correlation", *names])
    for i, name in enumerate(names):
        matrix.cell(offset+i+1, 1, name)
        for j, value in enumerate(d["correlation"][i], 2):
            if value is not None:
                matrix.cell(offset+i+1, j, value).number_format = RATIO
    matrix.conditional_formatting.add(
        f"B{offset+1}:{get_column_letter(len(names)+1)}{offset+len(names)}",
        ColorScaleRule(start_type="num", start_value=-1, start_color="E66A6A",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=1, end_color="4389E4"))
    if kind == "Markowitz Optimization":
        optimum = solution.diagnostics
        _header(results, 3, ["Investment", "Current weight", "Optimized weight", "Baseline weight"])
        for i, name in enumerate(names, 4):
            results.cell(i, 1, name)
            if d["current_weights"] is not None:
                results.cell(i, 2, d["current_weights"][name]).number_format = PERCENT
            results.cell(i, 3, optimum["weights"][name]).number_format = PERCENT
            if baseline is not None:
                results.cell(i, 4, baseline.diagnostics["weights"][name]).number_format = PERCENT
        r = len(names)+6
        for label, value, fmt in [
            ("Expected return per period", optimum["expected_return"], PERCENT),
            ("Variance (decimal return squared)", optimum["variance"], _variance_format(optimum["variance"])),
            ("Volatility per period", optimum["volatility"], PERCENT),
            ("Estimated Sharpe", optimum["estimated_sharpe"], RATIO),
            ("Solver status", optimum["solver_message"], None),
        ]:
            results.cell(r, 1, label)
            if value is not None:
                results.cell(r, 2, value)
                if fmt:
                    results.cell(r, 2).number_format = fmt
            r += 1
        results.cell(r+1, 1, "Constraint residuals (decimal)")
        for row, (label, value) in enumerate(optimum["constraint_residuals"].items(), r+2):
            results.cell(row, 1, label)
            if value is not None:
                results.cell(row, 2, value)
        if frontier is not None and frontier.ok:
            chart_ws = _sheet(book, "Frontier", "Achieved efficient frontier")
            _header(chart_ws, 3, ["Achieved volatility", "Achieved return", "Target floor"])
            for row, point in enumerate(frontier.diagnostics["points"], 4):
                for col, key in enumerate(("volatility", "return", "target"), 1):
                    chart_ws.cell(row, col, point[key]).number_format = PERCENT
            if len(frontier.diagnostics["points"]) > 1:
                chart = ScatterChart()
                chart.style = 10
                chart.title = "Efficient frontier"
                chart.x_axis.title = "Volatility (%)"
                chart.y_axis.title = "Expected return (%)"
                chart.x_axis.numFmt = "0.0%"
                chart.y_axis.numFmt = "0.0%"
                last = 3+len(frontier.diagnostics["points"])
                chart.series.append(Series(Reference(chart_ws, min_col=2, min_row=4, max_row=last),
                                           Reference(chart_ws, min_col=1, min_row=4, max_row=last), title="Frontier"))
                chart.series[0].graphicalProperties.line.solidFill = BLUE
                chart_ws.add_chart(chart, "E3")
    else:
        _header(results, 3, ["Portfolio", "Average per period", "Cumulative return", "Sample variance", "Volatility", "Sharpe"])
        for row, item in enumerate(d["comparison"], 4):
            for col, key in enumerate(("name", "mean", "cumulative", "variance", "volatility", "sharpe"), 1):
                if item[key] is not None:
                    results.cell(row, col, item[key])
                    if col in (2, 3, 5):
                        results.cell(row, col).number_format = PERCENT
                    elif col == 4:
                        results.cell(row, col).number_format = _variance_format(item[key])
                    elif col == 6:
                        results.cell(row, col).number_format = RATIO
        chart_ws = _sheet(book, "Performance", "Cumulative historical performance")
        _header(chart_ws, 3, ["Date", *names])
        for row, day in enumerate(inputs["dates"], 4):
            chart_ws.cell(row, 1, day)
            for col, item in enumerate(d["comparison"], 2):
                chart_ws.cell(row, col, item["path"][row-4]).number_format = PERCENT
        chart = LineChart()
        chart.style = 10
        chart.title = "Cumulative return"
        chart.y_axis.title = "Cumulative return (%)"
        chart.add_data(Reference(chart_ws, min_col=2, max_col=len(names)+1,
                                 min_row=3, max_row=3+len(inputs["dates"])), titles_from_data=True)
        for series, name in zip(chart.series, names):
            series.title = SeriesLabel(v=name)
        for index, series in enumerate(chart.series):
            series.graphicalProperties.line.solidFill = (BLUE, "168B50", "E68A2E", "8A63B8")[index % 4]
        chart.set_categories(Reference(chart_ws, min_col=1, min_row=4, max_row=3+len(inputs["dates"])))
        chart.y_axis.numFmt = "0%"
        chart_ws.add_chart(chart, "E3")
    original = inputs.get("original_import")
    if original:
        provenance = _sheet(book, "Original Import", "Original imported data and units")
        provenance["A3"], provenance["B3"] = "Source", original["source"]
        provenance["A4"], provenance["B4"] = "Original scale", original.get("original_scale", "Summary decimals")
        provenance["A5"], provenance["B5"] = "Meaning", original.get("purpose", "Individual investment estimates")
        if original["kind"] == "history":
            _header(provenance, 7, ["Date", *original["names"]])
            for row_no, (day, values) in enumerate(zip(original["dates"], original["original_values"]), 8):
                provenance.cell(row_no, 1, day)
                for col, value in enumerate(values, 2):
                    provenance.cell(row_no, col, value)
        else:
            _header(provenance, 7, ["Investment", "Arithmetic average (%)", "Current weight (%)"])
            for row_no, name in enumerate(original["names"], 8):
                provenance.cell(row_no, 1, name)
                provenance.cell(row_no, 2, original["expected_returns"][name]).number_format = PERCENT
                provenance.cell(row_no, 3, original["current_weights"][name]).number_format = PERCENT
            matrix_row = 10+len(original["names"])
            _header(provenance, matrix_row, ["Covariance (decimal return squared)", *original["names"]])
            for i, name in enumerate(original["names"]):
                provenance.cell(matrix_row+1+i, 1, name)
                for j, value in enumerate(original["covariance"][i], 2):
                    provenance.cell(matrix_row+1+i, j, value).number_format = _variance_format(value)
    dashboard = (_markowitz_dashboard(book, inputs, analysis, solution, frontier, baseline)
                 if kind == "Markowitz Optimization" else _compare_dashboard(book, inputs, analysis))
    _style_details(book, dashboard)
    for chart in dashboard._charts:
        chart.visible_cells_only = False
    for sheet in book:
        if sheet is not dashboard:
            sheet.sheet_state = "hidden"
    book.active = 0
    data = io.BytesIO()
    book.save(data)
    book.close()
    return data.getvalue()
