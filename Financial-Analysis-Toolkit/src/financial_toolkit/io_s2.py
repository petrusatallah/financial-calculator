"""S2 minimum typed CSV/Excel round trip. REQ-IO-01/REPORT-02; D12."""
from __future__ import annotations

import csv
import json
import math
import time
import zipfile
from pathlib import Path

from .bond_s2 import analyze_regular_bond
from .tvm_s2 import schedule_value, solve_regular
from .validation import InputError, parse_currency

class ResourceLimit(InputError):
    """The approved S2 time or estimated memory limit interrupted file work."""


def _deadline(started: float, mode: str = "Standard") -> None:
    limit = 30 if mode == "Advanced" else 10
    if time.monotonic() - started > limit:
        raise ResourceLimit(f"INCOMPLETE: {mode} file processing exceeded {limit} seconds.")

SCHEMA_VERSION = "1.0"
CALCULATIONS = {"tvm_regular": solve_regular, "tvm_schedule": schedule_value,
                "bond_regular": analyze_regular_bond}
RATE_KEYS = {"annual_rate", "coupon_rate", "growth_rate", "annual_yield", "bump"}
DATE_KEYS = {"valuation_date"}


def _plain(value):
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise InputError("Non-finite numeric value cannot be exported.")
        return value
    if isinstance(value, (list, tuple)):
        return [_plain(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _plain(item) for key, item in value.items()}
    if hasattr(value, "isoformat"):
        return value.isoformat()
    raise InputError(f"Unsupported file value type: {type(value).__name__}.")


def _input_meta(name: str, value, currency: str) -> dict:
    if name == "currency":
        return {"unit": "three-letter currency code", "scale": "identity"}
    if name in RATE_KEYS:
        return {"unit": "decimal rate", "scale": "decimal"}
    if name in DATE_KEYS:
        return {"unit": "ISO date", "scale": "calendar"}
    if name in {"face", "redemption", "observed_price", "pv", "fv", "pmt"}:
        return {"unit": currency, "scale": "one"}
    if name in {"periods", "payment_frequency", "compounding_frequency", "frequency"}:
        return {"unit": "count", "scale": "integer"}
    if name == "cash_flows":
        return {"unit": f"signed {currency} and time", "scale": "explicit"}
    return {"unit": "label or dimensionless", "scale": "explicit"}


def _normalize_inputs(kind: str, inputs: dict) -> dict:
    values = _plain(inputs)
    number_fields = {"pv", "fv", "pmt", "annual_rate", "coupon_rate", "growth_rate",
                     "face", "redemption", "annual_yield", "observed_price", "years", "bump"}
    integer_fields = {"periods", "payment_frequency", "compounding_frequency", "frequency"}
    for name in number_fields & values.keys():
        if values[name] is not None and values[name] != "":
            values[name] = float(values[name])
    for name in integer_fields & values.keys():
        if values[name] is not None and values[name] != "":
            number = float(values[name])
            if not number.is_integer():
                raise InputError(f"{name} must be a whole number.")
            values[name] = int(number)
    if "scenario_yields" in values:
        values["scenario_yields"] = [float(x) if not isinstance(x, dict) else
            {"label": str(x["label"]), "yield": float(x["yield"])} for x in values["scenario_yields"]]
    if "cash_flows" in values:
        dated = bool(values.get("dated"))
        rows = []
        for t, amount in values["cash_flows"]:
            if dated:
                time_value = str(t)
            else:
                number = float(t)
                if not number.is_integer():
                    raise InputError("Regular custom time must be a whole payment period; it cannot be rounded.")
                time_value = int(number)
            rows.append([time_value, float(amount)])
        values["cash_flows"] = rows
    return values

def make_record(calculation_type: str, inputs: dict) -> dict:
    if calculation_type not in CALCULATIONS:
        raise InputError("Unsupported calculation type for the S2 schema.")
    normalized = _normalize_inputs(calculation_type, inputs)
    normalized["currency"] = parse_currency(normalized.get("currency"))
    if not isinstance(normalized, dict):
        raise InputError("Inputs must be named fields.")
    result = CALCULATIONS[calculation_type](**normalized)
    if not result.ok:
        raise InputError(f"Calculation did not succeed: {result.status}: {result.message}")
    return {"schema_version": SCHEMA_VERSION, "calculation_type": calculation_type,
            "inputs": normalized,
            "input_metadata": {name: _input_meta(name, value, normalized["currency"]) for name, value in normalized.items()},
            "output": {"value": _plain(result.value), "unit": result.unit, "status": result.status.value,
                       "workings": _plain(result.diagnostics)},
            "assumptions": _plain(result.assumptions), "warnings": _plain(result.warnings)}


def recalculate_record(record: dict):
    required = {"schema_version", "calculation_type", "inputs", "input_metadata", "output", "assumptions", "warnings"}
    if not isinstance(record, dict) or set(record) != required:
        raise InputError("Analysis record fields do not match the S2 schema.")
    if not isinstance(record, dict) or record.get("schema_version") != SCHEMA_VERSION:
        raise InputError("Unsupported or missing schema version.")
    kind = record.get("calculation_type")
    if kind not in CALCULATIONS:
        raise InputError("Unknown calculation type.")
    inputs = record.get("inputs")
    metadata = record.get("input_metadata")
    if not isinstance(inputs, dict) or not isinstance(metadata, dict) or set(inputs) != set(metadata):
        raise InputError("Input names and metadata must match exactly.")
    currency = parse_currency(inputs.get("currency"))
    for name, value in inputs.items():
        if metadata[name] != _input_meta(name, value, currency):
            raise InputError(f"Incorrect unit or scale for {name}.")
    result = CALCULATIONS[kind](**inputs)
    if not result.ok:
        raise InputError(f"Imported inputs did not calculate: {result.status}: {result.message}")
    return result  # Imported stored output is never used as an oracle.


def _safe_text(value: str) -> str:
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def _restore_text(value: str) -> str:
    return value[1:] if value.startswith("'") and len(value) > 1 and value[1] in "=+-@\t\r" else value


def export_csv(path: str | Path, record: dict) -> None:
    recalculate_record(record)
    with open(path, "w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["section", "key", "json_value", "unit", "scale"])
        for key in ("schema_version", "calculation_type", "output", "assumptions", "warnings"):
            writer.writerow(["metadata", key, _safe_text(json.dumps(record[key], ensure_ascii=False)), "", ""])
        for name, value in record["inputs"].items():
            meta = record["input_metadata"][name]
            writer.writerow(["input", _safe_text(name), _safe_text(json.dumps(value, ensure_ascii=False)), meta["unit"], meta["scale"]])


def _file_guard(path: str | Path) -> None:
    # Cap both stored and expanded bytes before parsing. At most two bounded
    # worksheet tables are accepted; the conservative 16x estimate is 256 MiB.
    target = Path(path)
    if target.stat().st_size > 16 * 1024 * 1024:
        raise ResourceLimit("INCOMPLETE: typed file exceeds the 16 MiB S2 import bound.")
    if target.suffix.lower() == ".xlsx":
        with zipfile.ZipFile(target) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(item.file_size for item in entries) > 16 * 1024 * 1024:
                raise ResourceLimit("INCOMPLETE: workbook expansion exceeds the S2 memory bound.")

def import_csv(path: str | Path):
    started = time.monotonic()
    _file_guard(path)
    with open(path, "r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != ["section", "key", "json_value", "unit", "scale"]:
            raise InputError("CSV columns do not match the S2 schema.")
        record, inputs, meta = {}, {}, {}
        for row_no, row in enumerate(reader, 2):
            _deadline(started, inputs.get("mode", "Advanced"))
            if row_no > 10050:
                raise ResourceLimit("INCOMPLETE: typed CSV row count exceeds the S2 bound.")
            section, key = row["section"], _restore_text(row["key"])
            try:
                value = json.loads(_restore_text(row["json_value"]))
            except (TypeError, ValueError) as exc:
                raise InputError(f"Invalid JSON value at CSV row {row_no}.") from exc
            if section == "metadata":
                if key in record:
                    raise InputError(f"Duplicate CSV metadata field {key}.")
                record[key] = value
            elif section == "input":
                if key in inputs:
                    raise InputError(f"Duplicate CSV input field {key}.")
                inputs[key] = value
                meta[key] = {"unit": row["unit"], "scale": row["scale"]}
            else:
                raise InputError(f"Unknown CSV section at row {row_no}.")
    record["inputs"], record["input_metadata"] = inputs, meta
    result = recalculate_record(record)
    _deadline(started, inputs.get("mode", "Standard"))
    return record, result


def export_excel(path: str | Path, record: dict) -> None:
    from openpyxl import Workbook
    recalculate_record(record)
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "Metadata"
    metadata.append(["key", "json_value"])
    for key in ("schema_version", "calculation_type", "output", "assumptions", "warnings"):
        metadata.append([key, json.dumps(record[key], ensure_ascii=False)])
    sheet = workbook.create_sheet("Inputs")
    sheet.append(["key", "json_value", "unit", "scale"])
    for name, value in record["inputs"].items():
        meta = record["input_metadata"][name]
        sheet.append([name, json.dumps(value, ensure_ascii=False), meta["unit"], meta["scale"]])
    for ws in workbook:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Formula-looking text stays text.
    workbook.save(path)


def import_excel(path: str | Path):
    from openpyxl import load_workbook
    started = time.monotonic()
    _file_guard(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if set(workbook.sheetnames) != {"Metadata", "Inputs"}:
            raise InputError("Excel workbook must contain exactly Metadata and Inputs sheets.")
        tables = {}
        for name in ("Metadata", "Inputs"):
            rows = []
            for row_no, cells in enumerate(workbook[name].iter_rows(), 1):
                _deadline(started, "Advanced")
                if row_no > 10050:
                    raise ResourceLimit("INCOMPLETE: typed Excel row count exceeds the S2 bound.")
                if any(cell.data_type == "f" for cell in cells):
                    raise InputError("Formula cells are not valid typed input data.")
                rows.append(tuple(cell.value for cell in cells))
            tables[name] = rows
        metadata, inputs_rows = tables["Metadata"], tables["Inputs"]
        if not metadata or not inputs_rows or metadata[0] != ("key", "json_value") or inputs_rows[0] != ("key", "json_value", "unit", "scale"):
            raise InputError("Excel columns do not match the S2 schema.")
        record, inputs, meta = {}, {}, {}
        for row_no, (key, raw) in enumerate(metadata[1:], 2):
            _deadline(started, "Advanced")
            if key in record or not isinstance(raw, str):
                raise InputError(f"Invalid or duplicate Metadata row {row_no}.")
            try:
                record[key] = json.loads(raw)
            except ValueError as exc:
                raise InputError(f"Invalid JSON value in Metadata row {row_no}.") from exc
        for row_no, (key, raw, unit, scale) in enumerate(inputs_rows[1:], 2):
            _deadline(started, inputs.get("mode", "Advanced"))
            if key in inputs or not isinstance(raw, str):
                raise InputError(f"Invalid or duplicate Inputs row {row_no}.")
            try:
                inputs[key] = json.loads(raw)
            except ValueError as exc:
                raise InputError(f"Invalid JSON value in Inputs row {row_no}.") from exc
            meta[key] = {"unit": unit, "scale": scale}
        record["inputs"], record["input_metadata"] = inputs, meta
        result = recalculate_record(record)
        _deadline(started, inputs.get("mode", "Standard"))
        return record, result
    finally:
        workbook.close()
