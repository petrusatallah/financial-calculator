"""S5 minimum typed files. Imported stored outputs are never calculation inputs."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
import time

from .market_data_s5 import FIELDS, Incomplete, analyze_data, profile
from .validation import InputError

VERSION = '1.0'
KEYS = {'schema_version','calculation_type','inputs','input_metadata','output','assumptions','warnings'}
META = {'observations':'explicit S5 observation fields; rates decimal; dates ISO; prices per declared unit',
        'mode':'Standard or Advanced'}


def make_record(rows, mode='Standard'):
    result = analyze_data(rows, mode)
    if not result.ok:
        raise InputError(f'{result.status}: {result.message}')
    return {'schema_version':VERSION,'calculation_type':'market_data',
        'inputs':{'observations':result.diagnostics['observations'],'mode':mode},'input_metadata':dict(META),
        'output':{'row_count':result.diagnostics['row_count'],'latest':result.diagnostics['latest'],
                  'return_summary':{key:{'count':len(values),'last':values[-1] if values else None}
                                    for key,values in result.diagnostics['returns'].items()}},
        'assumptions':result.assumptions,'warnings':list(result.warnings)}


def recalculate(record):
    if not isinstance(record,dict) or set(record) != KEYS or record['schema_version'] != VERSION or record['calculation_type'] != 'market_data':
        raise InputError('Unsupported S5 typed record/schema version.')
    if record['input_metadata'] != META or not isinstance(record['inputs'],dict) or set(record['inputs']) != {'observations','mode'}:
        raise InputError('S5 input units/schema metadata mismatch.')
    result = analyze_data(record['inputs']['observations'],record['inputs']['mode'])
    if not result.ok:
        raise InputError(f'{result.status}: {result.message}')
    return result


def _json(text):
    def unique(pairs):
        result = {}
        for key,value in pairs:
            if key in result:
                raise InputError('Duplicate JSON record field.')
            result[key] = value
        return result
    try:
        return json.loads(text, object_pairs_hook=unique, parse_constant=lambda _: (_ for _ in ()).throw(InputError('Nonfinite JSON value.')))
    except (TypeError,ValueError,RecursionError) as exc:
        raise InputError('Invalid typed JSON file value.') from exc


def _dump(value):
    return json.dumps(value,ensure_ascii=False,allow_nan=False,separators=(',',':'))


def csv_bytes(record):
    started = time.monotonic()
    recalculate(record)
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(('section','key','json_value'))
    for key in sorted(KEYS-{'inputs'}):
        writer.writerow(('record',key,_dump(record[key])))
    writer.writerow(('input','mode',_dump(record['inputs']['mode'])))
    for i,row in enumerate(record['inputs']['observations']):
        writer.writerow(('observation',str(i),_dump(row)))
    data = stream.getvalue().encode('utf-8-sig')
    profile(record['inputs']['mode'],len(record['inputs']['observations']),0,started)
    return data


def excel_bytes(record):
    started = time.monotonic()
    from openpyxl import Workbook
    recalculate(record)
    workbook = Workbook()
    meta = workbook.active
    meta.title = 'Record'
    meta.append(('key','part','json_chunk'))
    for key in sorted(KEYS-{'inputs'}):
        raw = _dump(record[key])
        for part,offset in enumerate(range(0,len(raw),30000)):
            meta.append((key,part,raw[offset:offset+30000]))
    meta.append(('mode',0,_dump(record['inputs']['mode'])))
    sheet = workbook.create_sheet('Observations')
    sheet.append(FIELDS)
    for row in record['inputs']['observations']:
        sheet.append([_dump(row[key]) for key in FIELDS])
    for page in workbook:
        for row in page:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = 's'  # provider text cannot become a formula
                    if len(cell.value)>32767:
                        raise InputError('Excel cell size exceeded; use the typed CSV format.')
    stream = io.BytesIO()
    workbook.save(stream)
    profile(record['inputs']['mode'],len(record['inputs']['observations']),0,started)
    return stream.getvalue()


def _guard(data, suffix):
    if not isinstance(data,bytes) or len(data)>16*1024*1024:
        raise Incomplete('Typed file exceeds the 16 MiB stored bound.')
    if suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist())>100 or sum(x.file_size for x in archive.infolist())>24*1024*1024:
                    raise Incomplete('Typed workbook exceeds its expanded-size bound.')
        except zipfile.BadZipFile as exc:
            raise InputError('Invalid Excel workbook.') from exc


def _import_bytes(data, suffix):
    started = time.monotonic()
    suffix = suffix.lower()
    if suffix not in {'.csv','.xlsx'}:
        raise InputError('Only typed CSV and XLSX files are supported.')
    _guard(data,suffix)
    record, observations = {}, []
    if suffix == '.csv':
        reader = csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
        if reader.fieldnames != ['section','key','json_value']:
            raise InputError('Typed CSV columns do not match S5.')
        mode = None
        for index,row in enumerate(reader):
            profile('Advanced',len(observations),0,started)
            if None in row or any(v is None for v in row.values()):
                raise InputError('Malformed CSV row.')
            value = _json(row['json_value'])
            if row['section']=='record' and row['key'] in KEYS-{'inputs'}:
                if row['key'] in record:
                    raise InputError('Duplicate CSV metadata.')
                record[row['key']] = value
            elif row['section']=='input' and row['key']=='mode' and mode is None:
                mode = value
            elif row['section']=='observation' and row['key']==str(len(observations)):
                observations.append(value)
            else:
                raise InputError('Unknown, duplicate or out-of-order typed CSV field.')
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            if workbook.sheetnames != ['Record','Observations']:
                raise InputError('Typed workbook sheets mismatch.')
            for sheet in workbook:
                if sheet.max_row>20010 or sheet.max_column>len(FIELDS):
                    raise Incomplete('Workbook dimensions exceed S5 bounds.')
                for row in sheet:
                    if any(cell.data_type=='f' for cell in row):
                        raise InputError('Formulas are not accepted as typed observation values.')
            meta = list(workbook['Record'].values)
            if meta[0] != ('key','part','json_chunk'):
                raise InputError('Workbook metadata header mismatch.')
            mode = None
            chunks = {}
            for key,part,value in meta[1:]:
                if key not in (KEYS-{'inputs'})|{'mode'} or not isinstance(part,int) or part != len(chunks.get(key,[])) or not isinstance(value,str):
                    raise InputError('Duplicate/unknown workbook metadata.')
                chunks.setdefault(key,[]).append(value)
            for key,parts in chunks.items():
                if key=='mode':
                    mode = _json(''.join(parts))
                else:
                    record[key] = _json(''.join(parts))
            values = workbook['Observations'].values
            if tuple(next(values)) != FIELDS:
                raise InputError('Observation columns/units mismatch.')
            for row in values:
                profile('Advanced',len(observations),0,started)
                observations.append(dict(zip(FIELDS,(_json(cell) for cell in row))))
        finally:
            workbook.close()
    record['inputs'] = {'mode':mode,'observations':observations}
    result = recalculate(record)
    profile(mode,len(observations),len({(r['provider'],r['instrument']) for r in observations}),started)
    return record,result


def import_bytes(data, suffix):
    try:
        return _import_bytes(data,suffix)
    except (csv.Error,UnicodeError,zipfile.BadZipFile,KeyError,StopIteration,TypeError,ValueError) as exc:
        if isinstance(exc,InputError):
            raise
        raise InputError('Malformed typed data file.') from exc


def import_file(path):
    path = Path(path)
    if path.stat().st_size>16*1024*1024:
        raise Incomplete('Typed file exceeds stored-size bound.')
    return import_bytes(path.read_bytes(),path.suffix)


def portfolio_assets(rows, mode='Standard'):
    """Use validated price observations in the accepted S4 engine, preserving gaps."""
    result = analyze_data(rows,mode)
    if not result.ok:
        raise InputError(result.message)
    groups = {}
    for row in result.diagnostics['observations']:
        if row['asset_class']=='yield':
            raise InputError('Par yields are not asset prices for portfolio-return calculation.')
        key = row['provider']+'/'+row['instrument']
        if key not in groups:
            groups[key] = {'id':key,'currency':row['currency'],'frequency':row['frequency'],
                'kind':'price','adjustment':row['adjustment'],'observations':[],
                'provenance':[]}
        groups[key]['observations'].append((row['date'],row['value']))
        groups[key]['provenance'].append(row)
    for key,asset in groups.items():
        # S5 has already certified which adjacent observations represent one
        # period. S4 must not reconstruct a ratio across an excluded gap.
        asset['kind'] = 'returns'
        asset['adjustment'] = 'one-period returns from '+asset['adjustment']
        mapped = {r['date']:r['return'] for r in result.diagnostics['returns'][key]}
        asset['observations'] = [(day,mapped.get(day)) for day,_ in asset['observations'][1:]]
    return list(groups.values())
