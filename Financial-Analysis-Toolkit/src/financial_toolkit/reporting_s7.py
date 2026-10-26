"""Approved S7 v2 typed records, multi-sheet Excel and structured CSV bundles.

REQ-IO-01/REPORT-02/03. Exact editable typed leaves are authoritative;
cached outputs are compared but never used as calculation inputs.
"""
from __future__ import annotations
import copy
from contextlib import ExitStack
import csv
from dataclasses import asdict
import io
import inspect
import json
import math
import time
import uuid
import zipfile

from .workspace_s7 import FUNCTIONS,execute,plain
from .result import Status
from .validation import InputError

VERSION='2.0'
ENVELOPE={'schema_version','calculation_id','calculation_type','versions','inputs','input_metadata',
          'assumptions','sources','outputs','diagnostics','warnings','reproducibility'}
HEAD=['path','type','value','unit','scale','part']
SHEETS=('Metadata','Inputs','Assumptions','Results','Warnings','Sources','CashFlows','Observations','Diagnostics')
FILES={s:s.lower()+'.csv' for s in SHEETS};FILES['Metadata']='manifest.csv'
RATES={'annual_rate','annual_yield','coupon_rate','growth_rate','annual_discount_rate','annual_finance_rate',
 'annual_reinvestment_rate','cash_annual_yield','funding_annual_rate','risk_free_annual_rate',
 'market_expected_annual_return','spread','floor','cap','bump','yield','rate','target_return','effective_annual_rate','nominal_annual_rate','payment_period_rate','annual_effective_rate','periodic_rate'}
MONEY={'face','redemption','observed_price','pv','fv','pmt','amount','coupon','cash_flow','present_value','npv'}


class Limit(InputError):pass


class Guard:
    def __init__(self,mode='Standard',cancelled=lambda:False):
        if mode not in ('Standard','Advanced'):raise InputError('Unknown file execution profile.')
        self.mode=mode;self.cancelled=cancelled;self.started=time.monotonic()
        self.bytes=(16 if mode=='Standard' else 32)*1024**2
        self.rows=100000 if mode=='Standard' else 500000
        self.seconds=10 if mode=='Standard' else 30
        self.row_count=0
        self.memory=(256 if mode=='Standard' else 512)*1024**2
        self.estimated=0
    def check(self,size=0):
        if self.cancelled():raise Limit('INCOMPLETE: operation cancelled.')
        if time.monotonic()-self.started>self.seconds:raise Limit('INCOMPLETE: file/report deadline exceeded.')
        if size>self.bytes:raise Limit('INCOMPLETE: file/expanded-memory bound exceeded.')
    def row(self):
        self.row_count+=1
        if self.row_count>self.rows:raise Limit('INCOMPLETE: typed row bound exceeded.')
        self.check()
    def reserve(self,amount):
        self.check()
        if amount<0 or self.estimated+amount>self.memory:
            raise Limit('INCOMPLETE: aggregate parser/transform memory estimate exceeds profile.')
        self.estimated+=amount
    def parsed_row(self,row):
        # Includes retained strings/nodes, inverse reconstruction and validation buffers.
        self.reserve(2048+32*sum(len(str(v).encode('utf-8')) for v in row))


def escape(key):return str(key).replace('~','~0').replace('/','~1')
def pointer(path):return '/'+('/'.join(escape(k) for k in path))


def metadata(path,inputs):
    key=str(path[-1]) if path else ''
    currency=inputs.get('currency') or inputs.get('analysis_inputs',{}).get('currency') or inputs.get('base',{}).get('currency','declared currency')
    if len(path)>=2 and path[-2] in {'weights','lower_bounds','upper_bounds'}:return {'unit':'decimal allocation weight','scale':'one'}
    if len(path)>=2 and path[-2]=='expected_returns':return {'unit':'expected simple return per observation period','scale':'decimal'}
    if path and path[0] in {'values_x','values_y'}:
        driver=inputs.get('driver_x' if path[0]=='values_x' else 'driver_y','') or ''
        if 'rate' in driver:return {'unit':'annual effective rate','scale':'decimal'}
    if 'observations' in path and isinstance(path[-1],int):
        return {'unit':'ISO observation date' if path[-1]==0 else 'declared asset price/decimal return/distribution; kind/currency retained',
                'scale':'calendar' if path[-1]==0 else 'explicit numeric'}
    if key=='value':
        try:
            parent=inputs
            for k in path[:-1]:parent=parent[k]
            if isinstance(parent,dict) and 'unit' in parent:return {'unit':parent['unit'],'scale':'declared observation scale'}
        except (TypeError,KeyError,IndexError):pass
    if key in RATES:
        if key in {'annual_discount_rate','annual_finance_rate','annual_reinvestment_rate','risk_free_annual_rate','cash_annual_yield','funding_annual_rate','market_expected_annual_return'}:unit='annual effective rate'
        elif key in {'effective_annual_rate','annual_effective_rate'}:unit='annual effective rate'
        elif key in {'annual_rate','nominal_annual_rate'}:unit='nominal annual rate'
        elif key in {'annual_yield','yield','nominal_annual_yield'}:unit='nominal annual yield'
        elif key in {'payment_period_rate','periodic_rate'}:unit='rate per payment period'
        elif key=='coupon_rate':unit='annual coupon rate'
        elif key=='target_return':unit='return floor per observation period'
        elif key=='growth_rate':unit='growth per payment period'
        elif key=='rate':unit=('annual effective rate' if inputs.get('direction')=='effective_to_nominal' else 'nominal annual rate') if 'direction' in inputs else 'index fixing rate per declared tenor'
        elif key=='bump':unit='annual yield/rate shift'
        else:unit='annual contractual rate'
        return {'unit':unit,'scale':'decimal'}
    if key=='volatility':return {'unit':'decimal rate per square-root year','scale':'decimal'}
    if 'years_squared' in key:return {'unit':'years squared','scale':'one'}
    if 'duration_years' in key:return {'unit':'years','scale':'one'}
    if 'price' in key and isinstance(path[-1],str):return {'unit':currency+' per stated face or identified observation unit','scale':'one'}
    if key in MONEY:return {'unit':currency,'scale':'one'}
    if key in {'valuation_date','date','as_of','retrieved_at'}:return {'unit':'ISO date/time; declared source timezone where applicable','scale':'calendar'}
    if key in {'time','years','reset_years','payment_years'}:return {'unit':'years unless dated schedule','scale':'one'}
    if key in {'seed','period','periods','frequency','payment_frequency','compounding_frequency','training_paths','evaluation_paths','substeps'}:
        return {'unit':'count/seed or declared frequency label','scale':'integer or label'}
    # Structured rows preserve explicit inner units and their shape, not guessed rate scales.
    if 'cash_flows' in path and isinstance(path[-1],int):return {'unit':'explicit payment-period count or ISO cash-flow date' if path[-1]==0 else currency+' signed flow','scale':'explicit'}
    return {'unit':'declared value/label; context in inputs and assumptions','scale':'identity'}


def _nodes(obj,path=()):
    yield path,obj
    if isinstance(obj,dict):
        for k,v in obj.items():yield from _nodes(v,path+(k,))
    elif isinstance(obj,list):
        for i,v in enumerate(obj):yield from _nodes(v,path+(i,))


def input_metadata(inputs):
    return {pointer(path):metadata(path,inputs) for path,value in _nodes(inputs) if not isinstance(value,(dict,list))}


def _sources(inputs):
    out=[]
    for path,obj in _nodes(inputs):
        if isinstance(obj,dict) and ('source' in obj or 'provider' in obj):
            fields={k:v for k,v in obj.items() if k in ('id','source','provider','instrument','asset_class','currency','unit','date','valuation_date','as_of','retrieved_at','source_timezone','frequency','adjustment','warnings','tenor')}
            out.append({'input_path':pointer(path),**fields})
    return out


def make_record(kind,inputs,result=None):
    request=copy.deepcopy(plain(inputs))
    if kind not in FUNCTIONS:raise InputError('Unknown canonical calculation type.')
    guard=Guard(request.get('mode','Standard'));guard.check(len(json.dumps(request).encode())*2)
    result=execute(kind,request) if result is None else result
    if not result.ok:raise InputError(f'{result.status}: cannot export an unsuccessful calculation.')
    mc=result.diagnostics.get('monte_carlo',{})
    if 'seed' in mc and kind=='bond_option':request['seed']=mc['seed']
    # Market-data replay embeds validated observations; never store a credentialed fetch request.
    if kind=='market_data':request={'rows':plain(result.diagnostics['observations']),'mode':request.get('mode','Standard')}
    outputs={'status':str(result.status),'value':plain(result.value),'unit':result.unit,'message':result.message,
        'fields':[{'path':pointer(p),'status':'UNDEFINED' if v is None else 'VALID',
                   'unit':metadata(p,request)['unit'],'reason':'See metric applicability diagnostics' if v is None else ''}
                  for p,v in _nodes(result.diagnostics) if not isinstance(v,(dict,list))]}
    record={'schema_version':VERSION,'calculation_id':str(uuid.uuid4()),'calculation_type':kind,
        'versions':{'toolkit':'0.1.0','record':'2.0','model':result.assumptions.get('model_version','accepted S2–S6 engine')},
        'inputs':request,'input_metadata':input_metadata(request),'assumptions':plain(result.assumptions),
        'sources':_sources(request),'outputs':outputs,'diagnostics':plain(result.diagnostics),
        'warnings':list(result.warnings),'reproducibility':{k:plain(mc[k]) for k in ('seed','generator','numpy_version','stream_construction','cumulative_generated_paths') if k in mc}}
    validate(record);guard.check(len(json.dumps(record,ensure_ascii=False).encode()))
    return record


def validate(record):
    if not isinstance(record,dict) or set(record)!=ENVELOPE or record['schema_version']!=VERSION:
        raise InputError('Metadata: unsupported version or missing/unknown envelope field.')
    if not isinstance(record['calculation_type'],str) or record['calculation_type'] not in FUNCTIONS:raise InputError('Metadata: unsupported calculation type.')
    def shape(value,keys,location):
        if not isinstance(value,dict) or set(value)!=set(keys):raise InputError(location+': missing/unknown required fields or invalid object type.')
    def text(value,location,nonempty=False):
        if not isinstance(value,str) or nonempty and not value.strip():raise InputError(location+': explicit text required.')
    versions=record['versions'];shape(versions,{'toolkit','record','model'},'Metadata/versions')
    for key,value in versions.items():text(value,'Metadata/versions/'+key,True)
    if versions['record']!=VERSION or versions['toolkit'] not in {'0.1.0','legacy import'} or versions['model'] not in {'accepted S2–S6 engine','ho-lee-binomial-v1','preserved legacy settings'}:
        raise InputError('Metadata/versions: unsupported toolkit, record or model version.')
    for name in ('assumptions','diagnostics'):
        if not isinstance(record[name],dict):raise InputError(name+': object required.')
    if not isinstance(record['warnings'],list) or any(not isinstance(w,str) for w in record['warnings']):raise InputError('Warnings: text list required.')
    if not isinstance(record['sources'],list) or any(not isinstance(s,dict) for s in record['sources']):raise InputError('Sources: identified source objects required.')
    if not isinstance(record['inputs'],dict):raise InputError('Inputs: named fields required.')
    outputs=record['outputs'];shape(outputs,{'status','value','unit','message','fields'},'Results/outputs')
    legal={str(s) for s in Status}
    if not isinstance(outputs['status'],str) or outputs['status'] not in legal:raise InputError('Results/status: unsupported completion status.')
    text(outputs['unit'],'Results/unit',True);text(outputs['message'],'Results/message')
    def number(v):
        try:return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)
        except OverflowError:return False
    value=outputs['value']
    if value is not None and not number(value) and not (isinstance(value,(list,tuple)) and all(number(v) for v in value)):
        raise InputError('Results/value: finite number, numeric sequence or typed null required.')
    if not isinstance(outputs['fields'],list):raise InputError('Results/fields: typed field list required.')
    seen=set()
    for i,field in enumerate(outputs['fields']):
        location=f'Results/fields/{i}';shape(field,{'path','status','unit','reason'},location)
        for key in ('path','unit','reason'):text(field[key],location+'/'+key,key!='reason')
        if not field['path'].startswith('/') or field['path'] in seen or not isinstance(field['status'],str) or field['status'] not in legal:raise InputError(location+': duplicate path or invalid status/path.')
        seen.add(field['path'])
    replay=record['reproducibility'];mc=record['diagnostics'].get('monte_carlo')
    applicable=record['calculation_type']=='bond_option' and mc is not None
    if applicable and (not isinstance(mc,dict) or not mc):raise InputError('Metadata/reproducibility: malformed stored Monte Carlo snapshot.')
    if applicable:
        keys={'seed','generator','numpy_version','stream_construction','cumulative_generated_paths'}
        shape(replay,keys,'Metadata/reproducibility')
        from .monte_carlo_s6 import parse_seed
        if not isinstance(replay['seed'],int) or isinstance(replay['seed'],bool):raise InputError('Metadata/reproducibility/seed: integer required.')
        parse_seed(replay['seed'])
        for key in ('generator','numpy_version','stream_construction'):text(replay[key],'Metadata/reproducibility/'+key,True)
        count=replay['cumulative_generated_paths']
        if replay['generator']!='PCG64' or not isinstance(count,int) or isinstance(count,bool) or count<=0:
            raise InputError('Metadata/reproducibility: invalid generator or cumulative path count.')
        if not isinstance(mc,dict) or any(mc.get(key)!=replay[key] for key in keys):raise InputError('Metadata/reproducibility: inconsistent stored Monte Carlo snapshot.')
        # Edited Inputs remain authoritative; recorded seed describes the old result only.
    elif not isinstance(replay,dict) or replay:
        raise InputError('Metadata/reproducibility: no Monte Carlo fields apply to this calculation.')
    def fields(fn,values,location,top=False,skip=()):
        if not isinstance(values,dict):raise InputError(location+': named input fields required.')
        params=inspect.signature(fn).parameters
        forbidden={'cancel','cancelled','transport','now','execution_profile'}
        allowed=set(params)-forbidden-set(skip)
        if top:allowed.add('financial_mode')
        extras=set(values)-allowed
        missing={k for k,p in params.items() if p.default is inspect.Parameter.empty and k not in skip}-set(values)
        if extras or missing:raise InputError(location+f': unknown fields {sorted(extras)}; missing required fields {sorted(missing)}.')
    kind=record['calculation_type'];values=record['inputs']
    fields(FUNCTIONS[kind],values,'Inputs',True)
    from .appraisal_s3 import analyze_project
    if kind in {'sensitivity','scenarios'}:fields(analyze_project,values['base'],'Inputs/base')
    if kind=='project_comparison':
        for i,project in enumerate(values['projects']):
            if not isinstance(project,dict) or set(project)!={'name','inputs'}:raise InputError(f'Inputs/projects/{i}: name and inputs required.')
            fields(analyze_project,project['inputs'],f'Inputs/projects/{i}/inputs')
    if kind=='optimization':
        from .portfolio_s4 import analyze_portfolio
        from .optimization_s4 import optimize_portfolio
        fields(analyze_portfolio,values['analysis_inputs'],'Inputs/analysis_inputs')
        fields(optimize_portfolio,values['optimization_inputs'],'Inputs/optimization_inputs',skip=('analysis',))
    try:uuid.UUID(record['calculation_id'])
    except (ValueError,TypeError,AttributeError) as exc:raise InputError('Metadata: invalid calculation identity.') from exc
    if not isinstance(record['inputs'],dict) or record['input_metadata']!=input_metadata(record['inputs']):
        raise InputError('Inputs: declared unit/scale metadata conflicts with the canonical field contract.')
    if record['calculation_type']=='bond_option' and record['inputs'].get('monte_carlo',True) and record['inputs'].get('seed') is None:
        raise InputError('Inputs/seed: Monte Carlo replay requires its actual seed.')
    plain(record)
    if record['sources']!=_sources(record['inputs']):raise InputError('Sources: provenance conflicts with authoritative inputs.')


def recalculate(record,cancelled=lambda:False):
    validate(record);guard=Guard(record['inputs'].get('mode','Standard'),cancelled)
    result=execute(record['calculation_type'],record['inputs'],cancelled)
    guard.check()
    # Preserve and expose comparisons; freshness warnings may legitimately change.
    old={'value':record['outputs'].get('value'),'warnings':record['warnings'],'diagnostics':record['diagnostics']}
    now={'value':plain(result.value),'warnings':list(result.warnings),'diagnostics':plain(result.diagnostics)}
    return result,{'stored_status':record['outputs'].get('status'),'recalculated_status':str(result.status),
        'recorded_value':old['value'],'recalculated_value':now['value'],
        'recorded_unit':record['outputs'].get('unit'),'recalculated_unit':result.unit,
        'units_changed':record['outputs'].get('unit')!=result.unit,
        'assumptions_changed':record['assumptions']!=plain(result.assumptions),
        'recorded_seed':record['reproducibility'].get('seed'),
        'recalculated_seed':result.diagnostics.get('monte_carlo',{}).get('seed'),
        'reproducibility_changed':record['reproducibility']!={k:plain(result.diagnostics.get('monte_carlo',{})[k]) for k in ('seed','generator','numpy_version','stream_construction','cumulative_generated_paths') if k in result.diagnostics.get('monte_carlo',{})},
        'value_changed':old['value']!=now['value'],'warnings_changed':old['warnings']!=now['warnings'],
        'diagnostics_changed':old['diagnostics']!=now['diagnostics'],
        'meaning':'Recorded output is comparison evidence only; current recalculation is authoritative. Execution timing diagnostics may differ on replay.'}


def _table(path):
    root=path[0] if path else None
    if root=='inputs':
        if 'cash_flows' in path:return 'CashFlows'
        if 'observations' in path or 'rows' in path and len(path)>2:return 'Observations'
        return 'Inputs'
    if root=='diagnostics' and any(k in path for k in ('schedule','cash_flow_schedule','cash_flow_timeline')):return 'CashFlows'
    return {'assumptions':'Assumptions','outputs':'Results','warnings':'Warnings','sources':'Sources',
            'diagnostics':'Diagnostics'}.get(root,'Metadata')


def tables(record,guard):
    validate(record);out={s:[HEAD] for s in SHEETS}
    for path,value in _nodes(record):
        guard.check()
        if isinstance(value,dict):typ='object';raw=json.dumps(list(value),ensure_ascii=False)
        elif isinstance(value,list):typ='array';raw=str(len(value))
        elif value is None:typ='null';raw=''
        elif isinstance(value,bool):typ='boolean';raw=str(value).lower()
        elif isinstance(value,int):typ='integer';raw=str(value)
        elif isinstance(value,float):typ='number';raw=repr(float(value))
        else:typ='string';raw=value
        meta=record['input_metadata'].get(pointer(path[1:]),{}) if path and path[0]=='inputs' else {}
        table=_table(path)
        for part,start in enumerate(range(0,max(1,len(raw)),20000)):
            row=[pointer(path),typ,raw[start:start+20000],meta.get('unit',''),meta.get('scale',''),str(part)]
            guard.row();guard.parsed_row(row);out[table].append(row)
    manifest={'tables':{s:{'file':FILES[s],'rows':len(rows)-1,'columns':HEAD} for s,rows in out.items() if len(rows)>1}}
    row=['@manifest','manifest',json.dumps(manifest,separators=(',',':')),'','','0']
    guard.row();guard.parsed_row(row);out['Metadata'].append(row)
    return {s:rows for s,rows in out.items() if len(rows)>1}


def _protect(x):return "'"+x if x and x[0] in "'=+-@\t\r" else x
def _restore(x):return x[1:] if len(x)>1 and x[0]=="'" and x[1] in "'=+-@\t\r" else x


def csv_bundle(record,cancelled=lambda:False):
    guard=Guard(record['inputs'].get('mode','Standard'),cancelled);target=io.BytesIO();expanded=0
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,rows in tables(record,guard).items():
            guard.check();stream=io.StringIO(newline='');writer=csv.writer(stream)
            for row in rows:guard.check();writer.writerow([_protect(v) for v in row])
            data=stream.getvalue().encode('utf-8-sig');expanded+=len(data);guard.check(expanded)
            archive.writestr(FILES[name],data)
    guard.check(len(target.getvalue()));return target.getvalue()


def excel_bytes(record,cancelled=lambda:False):
    from openpyxl import Workbook
    guard=Guard(record['inputs'].get('mode','Standard'),cancelled)
    wb=Workbook();wb.remove(wb.active);target=io.BytesIO()
    for name,rows in tables(record,guard).items():
        ws=wb.create_sheet(name)
        for row in rows:
            guard.check();ws.append(row)
            for cell in ws[ws.max_row]:cell.data_type='s'
    wb.save(target);wb.close();data=target.getvalue();_archive_guard(data,guard)
    return data


def _archive_guard(data,guard,parsing=False):
    guard.check(len(data))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members=archive.infolist()
        if len(members)>32:raise Limit('INCOMPLETE: archive member count exceeds32.')
        if len({i.filename for i in members})!=len(members):raise InputError('Duplicate archive member.')
        for member in members:
            if member.flag_bits&1 or '..' in member.filename.split('/') or member.filename.startswith(('/','\\')) or '\\' in member.filename or ':' in member.filename:
                raise InputError('Encrypted or unsafe archive member.')
            if any(word in member.filename.lower() for word in ('vbaproject','externallinks','connections.xml')):
                raise InputError('Macros or external workbook connections are unsupported.')
        expanded=sum(i.file_size for i in members);guard.check(expanded)
        if parsing:guard.reserve(16*expanded)
    return members


def _json(text):
    def unique(pairs):
        obj={}
        for k,v in pairs:
            if k in obj:raise InputError('Duplicate object key.')
            obj[k]=v
        return obj
    return json.loads(text,object_pairs_hook=unique,parse_constant=lambda _:(_ for _ in ()).throw(InputError('Nonfinite JSON.')))


def _decode(tables_by_name,guard):
    nodes={};input_meta={};manifest=None
    for name,rows in tables_by_name.items():
        if next(rows,None)!=HEAD:raise InputError(f'{name}: missing/duplicate/malformed columns.')
        for rowno,row in enumerate(rows,2):
            guard.row();guard.parsed_row(row)
            if len(row)!=6 or not all(isinstance(v,str) for v in row):raise InputError(f'{name} row{rowno}: malformed typed row.')
            path,typ,value,unit,scale,part=row
            if path=='@manifest':
                if name!='Metadata' or manifest is not None or typ!='manifest' or part!='0':raise InputError('Metadata: duplicate/invalid manifest.')
                manifest=_json(value);continue
            key=path
            entry=nodes.setdefault(key,{'type':typ,'chunks':[],'table':name,'unit':unit,'scale':scale})
            if entry['type']!=typ or entry['table']!=name or entry['unit']!=unit or entry['scale']!=scale or part!=str(len(entry['chunks'])):
                raise InputError(f'{name} row{rowno} {path}: duplicate/conflicting/unordered chunks.')
            entry['chunks'].append(value)
    if manifest is None or set(manifest)!= {'tables'} or set(manifest['tables'])!=set(tables_by_name):raise InputError('Metadata: missing or conflicting table manifest.')
    used=set()
    def construct(path,depth=0):
        p=pointer(path)
        if depth>64 or p not in nodes:raise InputError(f'{p}: missing field or excessive nesting.')
        entry=nodes[p];used.add(p);typ=entry['type'];raw=''.join(entry['chunks'])
        if entry['table']!=_table(path):raise InputError(f'{p}: incorrect table assignment.')
        try:
            if typ=='object':
                keys=_json(raw)
                if not isinstance(keys,list) or not all(isinstance(k,str) for k in keys) or len(set(keys))!=len(keys):raise ValueError
                return {k:construct(path+(k,),depth+1) for k in keys}
            if typ=='array':
                n=int(raw)
                if str(n)!=raw or not 0<=n<=guard.rows:raise ValueError
                return [construct(path+(i,),depth+1) for i in range(n)]
            if path and path[0]=='inputs':input_meta[pointer(path[1:])]={'unit':entry['unit'],'scale':entry['scale']}
            if typ=='null' and raw=='':return None
            if typ=='boolean' and raw in ('true','false'):return raw=='true'
            if typ=='string':return raw
            if typ=='integer':
                v=int(raw)
                if str(v)!=raw:raise ValueError
                return v
            if typ=='number':
                if raw.endswith('%'):
                    if entry['scale']!='percent':raise InputError(f'{p}: percent suffix conflicts with declared scale.')
                    raw=raw[:-1]
                v=float(raw)
                if not math.isfinite(v):raise ValueError
                return v
        except (ValueError,TypeError) as exc:
            if isinstance(exc,InputError):raise
            raise InputError(f'{entry["table"]} {p}: invalid {typ} value.') from exc
        raise InputError(f'{p}: unsupported type/value {typ}.')
    record=construct(())
    if used!=set(nodes):raise InputError('Unknown undeclared typed field/row.')
    # Permit explicitly declared percent rates, converting before canonical validation.
    for path,value in _nodes(record['inputs']):
        if isinstance(value,(dict,list)):continue
        p=pointer(path);meta=input_meta[p];expected=metadata(path,record['inputs'])
        if meta['unit']!=expected['unit']:raise InputError(f'Inputs{p}: incorrect units.')
        if meta['scale']=='percent' and expected['scale']=='decimal':
            if not isinstance(value,(float,int)) or isinstance(value,bool):raise InputError(f'Inputs{p}: numeric percentage required.')
            parent=record['inputs']
            for k in path[:-1]:parent=parent[k]
            parent[path[-1]]=value/100
            input_meta[p]['scale']='decimal'
        elif meta['scale']!=expected['scale']:raise InputError(f'Inputs{p}: incorrect scale.')
    if record['input_metadata']!=input_meta:
        # Canonical metadata remains decimal when a human edits a rate to declared percent.
        if record['input_metadata']!=input_metadata(record['inputs']):raise InputError('Inputs metadata conflicts with record metadata.')
    for name,spec in manifest['tables'].items():
        count=sum(len(e['chunks']) for e in nodes.values() if e['table']==name)
        if spec!={'file':FILES[name],'rows':count,'columns':HEAD}:raise InputError(f'{name}: manifest row count/columns mismatch.')
    validate(record);return record


def read_bytes(data,suffix,mode='Standard',cancelled=lambda:False):
    """Parse only: caller explicitly confirms recalculation in a separate action."""
    guard=Guard(mode,cancelled);guard.check(len(data));guard.reserve(4*len(data))
    try:
        if suffix not in ('.zip','.xlsx','.csv'):raise InputError('Choose a supported workbook, CSV bundle or legacy typed CSV.')
        if suffix=='.zip':
            members=_archive_guard(data,guard,True)
            names={i.filename for i in members}
            if not names<=set(FILES.values()) or 'manifest.csv' not in names:raise InputError('Archive contains undeclared files.')
            with zipfile.ZipFile(io.BytesIO(data)) as archive,ExitStack() as stack:
                def csv_rows(file):
                    binary=stack.enter_context(archive.open(file))
                    text=stack.enter_context(io.TextIOWrapper(binary,encoding='utf-8-sig',newline=''))
                    for row in csv.reader(text):
                        guard.check();yield [_restore(v) for v in row]
                tables_by_name={name:csv_rows(file) for name,file in FILES.items() if file in names}
                record=_decode(tables_by_name,guard)
        elif suffix=='.xlsx':
            _archive_guard(data,guard,True)
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            try:
                if 'Metadata' not in wb.sheetnames or not set(wb.sheetnames)<=set(SHEETS):
                    return _legacy(data,suffix,guard)
                first=next(wb['Metadata'].iter_rows(values_only=True),())
                if list(first)!=HEAD:return _legacy(data,suffix,guard)
                def sheet_rows(ws):
                    for cells in ws:
                        guard.check()
                        if any(c.data_type=='f' for c in cells):raise InputError(f'{ws.title}: executable formula cell.')
                        yield ['' if c.value is None else str(c.value) for c in cells]
                tables_by_name={name:sheet_rows(wb[name]) for name in wb.sheetnames}
                record=_decode(tables_by_name,guard)
            finally:wb.close()
        else:return _legacy(data,suffix,guard)
        guard.check();return record,{'source_schema':VERSION,'preview_only':True}
    except (UnicodeError,json.JSONDecodeError,zipfile.BadZipFile,KeyError,TypeError,ValueError,OverflowError,csv.Error,RecursionError) as exc:
        if isinstance(exc,InputError):raise
        raise InputError('File has malformed encoding, structure or typed fields.') from exc


def _legacy(data,suffix,guard):
    """Known formats only. Legacy readers are bounded and recalculate on confirmation."""
    # Parsing is implemented in memory by inspecting each known documented envelope.
    from . import io_s6,io_s5,io_s2
    if suffix=='.csv':
        rows=csv.reader(io.StringIO(data.decode('utf-8-sig')))
        header=next(rows,None)
        if header is None:raise InputError('Empty typed CSV.')
        if header==io_s6.COLUMNS:
            def restored():
                yield header
                for row in rows:
                    guard.row();guard.parsed_row(row)
                    yield [io_s6._restore_csv(v) for v in row]
            record=io_s6._read_rows(restored(),guard.check);io_s6._validate(record)
            return _upgrade(record,'s6-1.0')
        if header==['section','key','json_value','unit','scale']:
            envelope={};inputs={};meta={}
            for no,row in enumerate(rows,2):
                guard.row();guard.parsed_row(row)
                if len(row)!=5:raise InputError(f'CSV row{no}: malformed legacy row.')
                section,key,value,unit,scale=row;key=io_s2._restore_text(key)
                dest=inputs if section=='input' else envelope if section=='metadata' else None
                if dest is None or key in dest:raise InputError(f'CSV row{no}: duplicate/unknown field.')
                dest[key]=_json(io_s2._restore_text(value))
                if section=='input':meta[key]={'unit':unit,'scale':scale}
            envelope.update(inputs=inputs,input_metadata=meta)
            _validate_legacy_s2(envelope);return _upgrade(envelope,'S2 minimum1.0')
        if header==['section','key','json_value']:
            envelope={};obs=[];mode=None
            for no,row in enumerate(rows,2):
                guard.row();guard.parsed_row(row)
                if len(row)!=3:raise InputError(f'CSV row{no}: malformed legacy row.')
                section,key,value=row
                if section=='record':
                    if key in envelope:raise InputError('Duplicate legacy metadata.')
                    envelope[key]=_json(value)
                elif section=='input' and key=='mode' and mode is None:mode=_json(value)
                elif section=='observation' and key==str(len(obs)):obs.append(_json(value))
                else:raise InputError('Unknown/duplicate legacy observation row.')
            envelope['inputs']={'observations':obs,'mode':mode}
            _validate_legacy_s5(envelope);return _upgrade(envelope,'S5 minimum1.0')
        raise InputError('Unknown legacy CSV schema; no unit guessing.')
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        tables={}
        for ws in wb:
            rows=[]
            for cells in ws:
                guard.row()
                if any(c.data_type=='f' for c in cells):raise InputError(f'{ws.title}: formula cells are invalid.')
                values=[c.value if c.value is not None else '' for c in cells]
                guard.parsed_row(values);rows.append(values)
            tables[ws.title]=rows
        if set(tables)=={'S6 typed record'}:
            record=io_s6._read_rows(iter(tables['S6 typed record']),guard.check);io_s6._validate(record)
            return _upgrade(record,'s6-1.0')
        if set(tables)=={'Metadata','Inputs'}:
            m=tables['Metadata'];ins=tables['Inputs']
            if m[0]!=['key','json_value'] or ins[0]!=['key','json_value','unit','scale']:raise InputError('Unknown legacy workbook columns.')
            envelope={};inputs={};meta={}
            for key,raw in m[1:]:
                if key in envelope:raise InputError('Duplicate legacy metadata.')
                envelope[key]=_json(raw)
            for key,raw,unit,scale in ins[1:]:
                if key in inputs:raise InputError('Duplicate legacy input.')
                inputs[key]=_json(raw);meta[key]={'unit':unit,'scale':scale}
            envelope.update(inputs=inputs,input_metadata=meta);_validate_legacy_s2(envelope)
            return _upgrade(envelope,'S2 minimum1.0')
        # S5 native workbook uses Record and Observations; inspected below by exact headers.
        if set(tables)=={'Record','Observations'}:
            r=tables['Record'];obs=tables['Observations']
            if r[0]!=['key','part','json_chunk']:raise InputError('Unknown S5 metadata columns.')
            envelope={}
            chunks={}
            for key,part,raw in r[1:]:
                field=chunks.setdefault(key,[])
                if part!=len(field):raise InputError('Duplicate/unordered S5 metadata.')
                field.append(raw)
            for key,parts in chunks.items():envelope[key]=_json(''.join(parts))
            from .market_data_s5 import FIELDS
            if obs[0]!=list(FIELDS):raise InputError('Unknown S5 observation columns.')
            values=[]
            for row in obs[1:]:
                item={key:(_json(value) if value!='' else None) for key,value in zip(FIELDS,row)};values.append(item)
            mode=envelope.pop('mode',None)
            envelope['inputs']={'observations':values,'mode':mode}
            _validate_legacy_s5(envelope);return _upgrade(envelope,'S5 minimum1.0')
        raise InputError('Unknown legacy workbook schema.')
    finally:wb.close()


def _validate_legacy_s2(record):
    from .io_s2 import CALCULATIONS,_input_meta
    if set(record)!= {'schema_version','calculation_type','inputs','input_metadata','output','assumptions','warnings'} or record['schema_version']!='1.0' or record['calculation_type'] not in CALCULATIONS:
        raise InputError('Unknown S2 legacy envelope.')
    inputs=record['inputs'];ccy=inputs.get('currency')
    if record['input_metadata']!={k:_input_meta(k,v,ccy) for k,v in inputs.items()}:raise InputError('S2 units/scales mismatch.')


def _validate_legacy_s5(record):
    from .io_s5 import KEYS,META
    if set(record)!=KEYS or record['schema_version']!='1.0' or record['calculation_type']!='market_data' or record['input_metadata']!=META:
        raise InputError('Unknown S5 legacy envelope.')


def _upgrade(old,source):
    kind=old['calculation_type'];inputs=copy.deepcopy(old['inputs'])
    if kind=='market_data':inputs={'rows':inputs['observations'],'mode':inputs['mode']}
    # No engine execution before the separate review/recalculate action.
    mc=old['output'].get('workings',{}).get('monte_carlo',{})
    record={'schema_version':VERSION,'calculation_id':str(uuid.uuid4()),'calculation_type':kind,
        'versions':{'toolkit':'legacy import','record':VERSION,'model':'preserved legacy settings'},
        'inputs':inputs,'input_metadata':input_metadata(inputs),'assumptions':old['assumptions'],
        'sources':_sources(inputs),'outputs':{'status':old['output'].get('status','VALID'),
            'value':old['output'].get('value'),'unit':old['output'].get('unit','legacy units'),'message':'','fields':[]},
        'diagnostics':old['output'].get('workings',{}),'warnings':old['warnings'],
        'reproducibility':{k:mc[k] for k in ('seed','generator','numpy_version','stream_construction','cumulative_generated_paths') if k in mc}}
    validate(record);return record,{'source_schema':source,'preview_only':True}
