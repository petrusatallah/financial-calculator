"""Fixed approved provider endpoints. Secrets stay in headers, never output or records."""
from __future__ import annotations

import csv
from datetime import date, datetime, timedelta, timezone
import io
from http.client import HTTPException
import json
import os
import re
import time
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlsplit
from urllib.request import Request, build_opener, HTTPRedirectHandler

from .market_data_s5 import (Incomplete, analyze_data, business_day, observation, profile,
                             resample, utc, validate)
from .result import CalculationResult, Status
from .validation import InputError, UnsupportedInput

TENORS = ('1 Mo','3 Mo','6 Mo','1 Yr','2 Yr','3 Yr','5 Yr','7 Yr','10 Yr','20 Yr','30 Yr')
WB_PAGE = 'https://www.worldbank.org/en/research/commodity-markets'
WB_HOST = 'thedocs.worldbank.org'
EXCHANGES = {'NYSE', 'NASDAQ', 'NYSE ARCA', 'NYSEARCA', 'AMEX', 'BATS', 'NYSE MKT'}


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        # A redirect must never forward a personal API token to another host.
        raise Incomplete('Provider redirect rejected; approved endpoint requires review.')


def _token():
    token = os.environ.get('TIINGO_API_KEY')
    if not token and os.name == 'nt':
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, 'Environment') as key:
                token = winreg.QueryValueEx(key, 'TIINGO_API_KEY')[0]
        except OSError:
            pass
    if not token:
        raise Incomplete('Tiingo requires a personal TIINGO_API_KEY; configure it locally, never in exported records.')
    if not isinstance(token, str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,256}', token):
        raise InputError('Configured Tiingo token has an invalid format.')
    return token


def _download(url, started, mode, *, headers=None, transport=None, cancelled=None):
    if cancelled and cancelled():
        raise Incomplete('Data retrieval cancelled.')
    profile(mode, 0, 0, started)
    cap = (8 if mode == 'Standard' else 24) * 1024 * 1024
    remaining = (10 if mode == 'Standard' else 30) - (time.monotonic()-started)
    if remaining <= 0:
        raise Incomplete('Data retrieval deadline exceeded.')
    if transport:
        data = transport(url, headers or {}, remaining)
    else:
        request = Request(url, headers={'User-Agent': 'FinancialToolkit/1.0', **(headers or {})})
        with build_opener(NoRedirect).open(request, timeout=remaining) as response:
            length = response.headers.get('Content-Length')
            if length and int(length) > cap:
                raise Incomplete('Provider response exceeds bounded download size.')
            chunks, count = [], 0
            while True:
                if cancelled and cancelled():
                    raise Incomplete('Data retrieval cancelled.')
                profile(mode, 0, 0, started)
                chunk = response.read(min(65536, cap + 1-count))
                if not chunk:
                    break
                chunks.append(chunk)
                count += len(chunk)
                if count > cap:
                    raise Incomplete('Provider response exceeds bounded download size.')
            data = b''.join(chunks)
            if length and len(data) != int(length):
                raise Incomplete('Provider response was truncated.')
    if not isinstance(data, bytes) or len(data) > cap:
        raise Incomplete('Provider response exceeds bounded download size.')
    profile(mode, 0, 0, started)
    if cancelled and cancelled():
        raise Incomplete('Data retrieval cancelled.')
    return data


def _row(provider, instrument, kind, day, value, currency, unit, retrieved, source, tz, **kw):
    return observation(provider=provider, instrument=instrument, asset_class=kind, day=day,
        value=value, currency=currency, unit=unit, retrieved_at=retrieved, source=source,
        source_timezone=tz, **kw)


def parse_ecb(data, retrieved, source):
    rows = list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
    if not rows:
        raise Incomplete('ECB returned no observations.')
    result = []
    for r in rows:
        if (r.get('CURRENCY'), r.get('CURRENCY_DENOM'), r.get('FREQ'), r.get('EXR_TYPE'), r.get('EXR_SUFFIX')) != ('USD','EUR','D','SP00','A'):
            raise InputError('ECB series identity or frequency mismatch.')
        value = r.get('OBS_VALUE')
        result.append(_row('ecb','EURUSD','fx',r['TIME_PERIOD'],None if value in ('', None) else value,
            'USD','USD per EUR',retrieved,source,'Europe/Berlin',
            warnings=['ECB informational reference rate; not an executable quote.']))
    return result


def parse_treasury(data, retrieved, source):
    reader = csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
    if not reader.fieldnames or not set(('Date', *TENORS)).issubset(reader.fieldnames):
        raise InputError('Treasury CSV tenor schema mismatch.')
    result = []
    for r in reader:
        day = datetime.strptime(r['Date'], '%m/%d/%Y').date().isoformat()
        for tenor in TENORS:
            value = r[tenor]
            result.append(_row('treasury','UST:'+tenor,'yield',day,
                None if value in ('', 'N/A', None) else float(value)/100,
                'USD','decimal nominal annual par yield',retrieved,source,'America/New_York',
                warnings=['Indicative par yield; not a bond price, zero rate or investment return.']))
    if not result:
        raise Incomplete('Treasury returned no observations.')
    return result


def parse_tiingo(data, retrieved, source, instrument, kind, metadata):
    if not isinstance(metadata,dict) or metadata.get('ticker', '').upper() != instrument or str(metadata.get('exchangeCode','')).upper() not in EXCHANGES:
        raise UnsupportedInput('Tiingo symbol must resolve to an approved U.S. equity exchange.')
    raw = json.loads(data)
    if not isinstance(raw, list) or not raw:
        raise Incomplete('Tiingo returned no observations.')
    result = []
    for r in raw:
        if not isinstance(r, dict) or not {'date','close','adjClose','divCash','splitFactor'}.issubset(r):
            raise InputError('Tiingo prices lack required corporate-action fields.')
        day = utc(r['date']).date().isoformat()
        result.append(_row('tiingo',instrument,kind,day,r['adjClose'],'USD','USD per share',retrieved,
            source,'America/New_York',adjustment='split_and_dividend_adjusted_close',
            raw_close=r['close'],distribution=r['divCash'],split_factor=r['splitFactor'],
            source_timestamp=r['date'],
            warnings=['Personal Tiingo token-holder internal use only; no redistribution.',
                      'Adjusted close includes distributions; do not add divCash again.',
                      'U.S. exchange: '+str(metadata['exchangeCode'])]))
    return result


def parse_world_bank(data, retrieved, source, instrument):
    import zipfile
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 100 or sum(x.file_size for x in archive.infolist()) > 24*1024*1024:
            raise Incomplete('World Bank workbook expanded size exceeds data limit.')
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook['Monthly Prices']
        if sheet.max_row > 10000 or sheet.max_column > 256:
            raise Incomplete('World Bank workbook dimensions exceed data limit.')
        header = next(sheet.iter_rows(min_row=5,max_row=5,values_only=True))
        units = next(sheet.iter_rows(min_row=6,max_row=6,values_only=True))
        title, unit, label = {'BRENT': ('Crude oil, Brent', '($/bbl)', 'USD per barrel'),
                              'GOLD': ('Gold', '($/troy oz)', 'USD per troy ounce')}[instrument]
        index = header.index(title)
        if units[index] != unit:
            raise InputError('World Bank commodity unit mismatch.')
        updated = str(sheet['A4'].value)
        try:
            published = datetime.strptime(updated, 'Updated on %B %d, %Y').date().isoformat()
        except ValueError as exc:
            raise InputError('World Bank publication date schema changed.') from exc
        result = []
        for values in sheet.iter_rows(min_row=7,values_only=True):
            if not values[0]:
                continue
            if not re.fullmatch(r'\d{4}M\d{2}', str(values[0])):
                raise InputError('World Bank monthly date schema changed.')
            y, m = int(values[0][:4]), int(values[0][5:])
            import calendar
            day = date(y,m,calendar.monthrange(y,m)[1]).isoformat()
            value = values[index]
            if isinstance(value, str) and value.strip() in {'', '..', '…', '�'}:
                value = None
            result.append(_row('world_bank',instrument,'commodity',day,value,'USD',label,retrieved,
                source,'UTC',frequency='monthly',source_period=f'{y:04}-{m:02}',
                published_date=published,
                warnings=['World Bank Pink Sheet; CC BY 4.0. Nominal USD monthly average; not a spot quote.',
                          updated+'; publication clock unspecified.']))
        return result
    finally:
        workbook.close()


def _fill_missing(rows, start, end):
    """Retain explicit holes on expected publication sessions so returns cannot bridge them."""
    if not rows or rows[0]['frequency'] != 'daily':
        return rows
    groups = {}
    for row in rows:
        groups.setdefault(row['instrument'], []).append(row)
    result = list(rows)
    for values in groups.values():
        known = {r['date'] for r in values}
        # Do not infer pre-inception observations or unfinished latest sessions.
        first = max(start, date.fromisoformat(min(known)))
        last = min(end, date.fromisoformat(max(known)))
        cal = {'tiingo':'NYSE','treasury':'US_FEDERAL','ecb':'TARGET'}[values[0]['provider']]
        day = first
        while day <= last:
            if business_day(day, cal) and day.isoformat() not in known:
                out = dict(values[0], date=day.isoformat(), value=None, raw_close=None, distribution=None, split_factor=None)
                out['source_timestamp'] = day.isoformat()
                out['warnings'] = ['Expected publication session missing; no return across this gap.']
                result.append(out)
            day += timedelta(days=1)
    return result


def fetch_data(*, provider, instrument, start, end, frequency='daily', asset_class='stock',
               mode='Standard', transport=None, now=None, cancelled=None):
    started = time.monotonic()
    try:
        profile(mode,0,0,started)
        start, end = date.fromisoformat(str(start)), date.fromisoformat(str(end))
        retrieved = now or datetime.now(timezone.utc).isoformat()
        # A user's local "today" may be tomorrow in UTC. Request bounds may
        # include that date; returned observations still cannot postdate their
        # retrieval date in the declared source timezone (validate enforces it).
        if start > end or end > utc(retrieved).date()+timedelta(days=1):
            raise InputError('Dates must be ordered and cannot request distant future observations.')
        if (end-start).days > (365*10 if mode == 'Standard' else 365*40):
            raise UnsupportedInput('Requested history exceeds bounded S5 date span.')
        instrument = str(instrument).strip().upper()
        if provider == 'ecb':
            if instrument != 'EURUSD':
                raise UnsupportedInput('Approved ECB coverage is EURUSD, quoted USD per EUR.')
            url = 'https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A?'+urlencode(
                {'startPeriod':start.isoformat(),'endPeriod':end.isoformat(),'format':'csvdata'})
            rows = parse_ecb(_download(url,started,mode,transport=transport,cancelled=cancelled),retrieved,url)
        elif provider == 'treasury':
            if instrument != 'UST':
                raise UnsupportedInput('Approved Treasury request identity is UST (selected nominal par tenors).')
            months = []
            cursor = start.replace(day=1)
            while cursor <= end:
                months.append(cursor.strftime('%Y%m'))
                cursor = date(cursor.year+(cursor.month==12), cursor.month%12+1, 1)
            if len(months) > (12 if mode == 'Standard' else 48):
                raise UnsupportedInput('Treasury retrieval is bounded to 12/48 monthly requests.')
            rows = []
            for month in months:
                url = f'https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/all/{month}?'+urlencode(
                    {'_format':'csv','field_tdr_date_value_month':month,'page':'','type':'daily_treasury_yield_curve'})
                rows.extend(parse_treasury(_download(url,started,mode,transport=transport,cancelled=cancelled),retrieved,url))
                profile(mode,len(rows),len(TENORS),started)
        elif provider == 'tiingo':
            if asset_class not in {'stock','etf'} or not re.fullmatch(r'[A-Z][A-Z0-9.-]{0,14}',instrument):
                raise UnsupportedInput('Tiingo accepts approved U.S. stock/ETF symbols only.')
            headers = {'Authorization':'Token '+_token(), 'Accept':'application/json'}
            base = 'https://api.tiingo.com/tiingo/daily/'+instrument
            metadata = json.loads(_download(base,started,mode,headers=headers,transport=transport,cancelled=cancelled))
            url = base+'/prices?'+urlencode({'startDate':start.isoformat(),'endDate':end.isoformat(),'resampleFreq':'daily'})
            rows = parse_tiingo(_download(url,started,mode,headers=headers,transport=transport,cancelled=cancelled),
                retrieved,url,instrument,asset_class,metadata)
        elif provider == 'world_bank':
            if instrument not in {'BRENT','GOLD'} or frequency != 'monthly':
                raise UnsupportedInput('World Bank approved rows are monthly BRENT and GOLD.')
            page = _download(WB_PAGE,started,mode,transport=transport,cancelled=cancelled).decode('utf-8')
            links = re.findall(r'https://thedocs\.worldbank\.org/[^\s\"<>]+/CMO-Historical-Data-Monthly\.xlsx',page)
            links = list(dict.fromkeys(links))
            if len(links) != 1 or urlsplit(links[0]).hostname != WB_HOST:
                raise Incomplete('Current World Bank monthly link could not be resolved unambiguously.')
            url = links[0]
            rows = parse_world_bank(_download(url,started,mode,transport=transport,cancelled=cancelled),retrieved,url,instrument)
        else:
            raise UnsupportedInput('Provider is outside approved S5 coverage.')
        rows = [r for r in rows if start.isoformat() <= r['date'] <= end.isoformat()]
        if not rows:
            raise Incomplete('No observations exist in the requested window.')
        rows = _fill_missing(rows,start,end)
        rows = validate(rows,mode,started=started)
        missing = [{'provider':r['provider'],'instrument':r['instrument'],'date':r['date']} for r in rows if r['value'] is None]
        rows = resample(rows,frequency,mode)
        result = analyze_data(rows,mode,now=retrieved)
        profile(mode,len(rows),len({r['instrument'] for r in rows}),started)
        if cancelled and cancelled():
            raise Incomplete('Data retrieval cancelled.')
        if missing:
            return CalculationResult(Status.INCOMPLETE,None,'market data',message='Provider observations are missing; incomplete retrieval is not success.',
                diagnostics={**result.diagnostics,'missing_source_rows':missing},
                warnings=tuple(dict.fromkeys((*result.warnings,f'Underlying source history is incomplete: {len(missing)} missing observations.'))),
                assumptions=result.assumptions)
        return result
    except HTTPError as exc:
        msg = 'Provider rate limit reached; retry later.' if exc.code == 429 else f'Provider HTTP {exc.code}; retrieval unavailable.'
        return CalculationResult(Status.INCOMPLETE,None,'market data',message=msg)
    except (URLError, OSError, TimeoutError, HTTPException):
        return CalculationResult(Status.INCOMPLETE,None,'market data',message='Provider connection unavailable; manual/CSV/Excel calculations remain available.')
    except (InputError, TypeError, ValueError, KeyError, IndexError, OverflowError) as exc:
        status = Status.INCOMPLETE if isinstance(exc, Incomplete) else Status.UNSUPPORTED if isinstance(exc, UnsupportedInput) else Status.INVALID
        # Never expose a provider exception body/URL or an authentication header.
        msg = str(exc) if isinstance(exc, InputError) else 'Provider response failed its expected schema.'
        return CalculationResult(status,None,'market data',message=msg)
    except Exception:
        return CalculationResult(Status.INVALID,None,'market data',message='Provider response could not be safely parsed; no observation was accepted.')

