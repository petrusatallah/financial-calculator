"""S6 typed calculation records, bounded CSV/Excel and recalculating import.

REQ-REPORT-02, REQ-IO-01, REQ-MC-01. Each JSON field is chunked below
Excel's cell text limit. Stored prices/policies are evidence, never input oracles.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import time
import zipfile

from .bond_s6 import analyze_floating_bond, analyze_option_bond
from .io_s2 import _plain
from .validation import InputError

VERSION = 's6-1.0'
FUNCTIONS = {'bond_floating':analyze_floating_bond,'bond_option':analyze_option_bond}
COLUMNS = ['section','key','part','json_text','unit','scale']
MAX_BYTES = 16*1024**2


def metadata(key,ccy):
    if key in {'face','redemption','observed_price'}: return {'unit':ccy,'scale':'one'}
    if key in {'coupon_rate','spread','floor','cap'}: return {'unit':'annual decimal rate','scale':'decimal'}
    if key=='volatility': return {'unit':'decimal rate per square-root year','scale':'decimal'}
    if key in {'frequency','substeps','training_paths','evaluation_paths','seed'}:
        return {'unit':'count or PRNG seed integer','scale':'integer'}
    if key=='years': return {'unit':'years','scale':'one'}
    if key=='valuation_date': return {'unit':'ISO calendar date','scale':'calendar'}
    if key in {'projection_curve','discount_curve'}:
        return {'unit':'identified DF curve: times in years, dimensionless positive DF, ISO as-of with timezone','scale':'explicit nested'}
    if key=='fixings': return {'unit':'coupon period; decimal index rate; source; ISO as-of with timezone','scale':'explicit nested'}
    if key=='exercises': return {'unit':f'coupon period; call/put {ccy}; contractual priority','scale':'explicit nested'}
    if key=='scenario_yields': return {'unit':f'label; decimal nominal annual yield; coupon period; redemption {ccy}','scale':'explicit nested'}
    return {'unit':'explicit label or boolean','scale':'identity'}


def make_record(kind,inputs):
    if kind not in FUNCTIONS or not isinstance(inputs,dict): raise InputError('Unknown S6 calculation request.')
    normalized = copy.deepcopy(_plain(inputs))
    if 'cancelled' in normalized: raise InputError('Cancellation callbacks cannot be serialized.')
    if kind=='bond_option' and normalized.get('monte_carlo',True):
        from .monte_carlo_s6 import parse_seed
        normalized['seed'] = parse_seed(normalized.get('seed'))
    result = FUNCTIONS[kind](**normalized)
    if not result.ok: raise InputError(f'{result.status}: {result.message}')
    return record_from_result(kind,normalized,result)


def record_from_result(kind,inputs,result):
    """Capture the UI's completed calculation, including the actual random seed."""
    if kind not in FUNCTIONS or not result.ok: raise InputError('Only completed S6 calculations can be exported.')
    normalized=copy.deepcopy(_plain(inputs))
    if kind=='bond_option' and normalized.get('monte_carlo',True):
        normalized['seed']=result.diagnostics['monte_carlo']['seed']
    ccy = result.assumptions['currency']; normalized['currency']=ccy
    return {'schema_version':VERSION,'calculation_type':kind,'inputs':normalized,
        'input_metadata':{k:metadata(k,ccy) for k in normalized},
        'output':{'status':str(result.status),'value':result.value,'unit':result.unit,'workings':_plain(result.diagnostics)},
        'assumptions':_plain(result.assumptions),'warnings':list(result.warnings)}


def _validate(record):
    required={'schema_version','calculation_type','inputs','input_metadata','output','assumptions','warnings'}
    if not isinstance(record,dict) or set(record)!=required or record['schema_version']!=VERSION:
        raise InputError('Invalid S6 schema or envelope.')
    if record['calculation_type'] not in FUNCTIONS: raise InputError('Unknown S6 calculation type.')
    inputs=record['inputs']; meta=record['input_metadata']
    if not isinstance(inputs,dict) or not isinstance(meta,dict) or set(inputs)!=set(meta):
        raise InputError('S6 input metadata does not match fields.')
    from .validation import parse_currency
    ccy=parse_currency(inputs.get('currency'))
    for k in inputs:
        if meta[k]!=metadata(k,ccy): raise InputError(f'Incorrect unit/scale for {k}.')
    if record['calculation_type']=='bond_option' and inputs.get('monte_carlo',True) and inputs.get('seed') is None:
        raise InputError('Monte Carlo records require the actual seed for replay.')
    _plain(record)
    if 'cancelled' in inputs: raise InputError('Cancellation is not a stored input field.')


def recalculate_record(record):
    _validate(record)
    try: result=FUNCTIONS[record['calculation_type']](**record['inputs'])
    except TypeError as exc: raise InputError('Unknown/missing S6 input field.') from exc
    if not result.ok: raise InputError(f'Recalculated inputs: {result.status}: {result.message}')
    return result


def _rows(record):
    _validate(record)
    rows=[COLUMNS]
    for key,value in record.items():
        if key in {'inputs','input_metadata'}: continue
        raw=json.dumps(value,ensure_ascii=False,allow_nan=False)
        for part,start in enumerate(range(0,len(raw),24000)):
            rows.append(['metadata',key,part,raw[start:start+24000],'',''])
    for key,value in record['inputs'].items():
        raw=json.dumps(value,ensure_ascii=False,allow_nan=False)
        meta=record['input_metadata'][key]
        for part,start in enumerate(range(0,len(raw),24000)):
            rows.append(['input',key,part,raw[start:start+24000],meta['unit'],meta['scale']])
    return rows


def csv_bytes(record):
    stream=io.StringIO(newline=''); writer=csv.writer(stream)
    for row in _rows(record):
        writer.writerow([_protect_csv(x) for x in row])
    return stream.getvalue().encode('utf-8-sig')


def _protect_csv(x):
    return "'"+x if isinstance(x,str) and x and x[0] in "'=+-@\t\r" else x


def _restore_csv(x):
    return x[1:] if x.startswith("'") and len(x)>1 and x[1] in "'=+-@\t\r" else x


def excel_bytes(record):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title='S6 typed record'
    for row in _rows(record): ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
    output=io.BytesIO(); wb.save(output); wb.close()
    if output.tell()>MAX_BYTES: raise InputError('S6 export exceeds file size limit.')
    return output.getvalue()


def _unique_object(pairs):
    obj={}
    for k,v in pairs:
        if k in obj: raise InputError('Duplicate JSON object key.')
        obj[k]=v
    return obj


def import_bytes(data,suffix,cancelled=lambda:False):
    started=time.monotonic()
    def guard():
        if cancelled(): raise InputError('INCOMPLETE: import cancelled.')
        if time.monotonic()-started>30: raise InputError('INCOMPLETE: import time bound exceeded.')
    guard()
    if len(data)>MAX_BYTES: raise InputError('INCOMPLETE: file exceeds 16 MiB bound.')
    if suffix=='.csv':
        rows=([_restore_csv(x) for x in row] for row in csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        record=_read_rows(rows,guard)
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist())>100 or sum(item.file_size for item in archive.infolist())>MAX_BYTES:
                raise InputError('INCOMPLETE: expanded workbook exceeds bounds.')
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            if wb.sheetnames!=['S6 typed record']: raise InputError('Invalid S6 workbook sheets.')
            def rows():
                for cells in wb.active:
                    if any(c.data_type=='f' for c in cells): raise InputError('Formula cells are not typed data.')
                    yield [c.value if c.value is not None else '' for c in cells]
            record=_read_rows(rows(),guard)
        finally: wb.close()
    else: raise InputError('Choose .csv or .xlsx.')
    guard()
    _validate(record)
    result=FUNCTIONS[record['calculation_type']](**record['inputs'],cancelled=cancelled)
    guard()
    limit=10 if record['inputs'].get('mode','Standard')=='Standard' else 30
    if time.monotonic()-started>limit: raise InputError('INCOMPLETE: profile import time bound exceeded.')
    if not result.ok: raise InputError(f'Recalculated inputs: {result.status}: {result.message}')
    return record,result


def _read_rows(rows,guard):
    rows=iter(rows)
    if list(next(rows,[]))!=COLUMNS: raise InputError('Invalid S6 typed columns.')
    fields={}; metas={}
    for rowno,row in enumerate(rows,1):
        guard()
        if rowno>20000 or len(row)!=6: raise InputError('Invalid or excessive S6 rows.')
        section,key,part,text,unit,scale=row
        if section not in {'metadata','input'} or not isinstance(key,str) or not isinstance(text,str):
            raise InputError('Invalid S6 record field.')
        if section=='metadata' and (unit or scale): raise InputError('Unexpected metadata units.')
        try:
            index=int(part)
            if str(index)!=str(part): raise ValueError
        except (TypeError,ValueError) as exc: raise InputError('Invalid S6 chunk number.') from exc
        field=fields.setdefault((section,key),[])
        if index!=len(field): raise InputError('Duplicate, missing or unordered field chunks.')
        field.append(text)
        if section=='input':
            meta={'unit':unit,'scale':scale}
            if key in metas and metas[key]!=meta: raise InputError('Conflicting chunk metadata.')
            metas[key]=meta
    record={}; inputs={}
    for (section,key),chunks in fields.items():
        try: value=json.loads(''.join(chunks),object_pairs_hook=_unique_object,
                              parse_constant=lambda x:(_ for _ in ()).throw(InputError('Nonfinite JSON.')))
        except ValueError as exc: raise InputError('Malformed field JSON.') from exc
        (inputs if section=='input' else record)[key]=value
    if 'inputs' in record or 'input_metadata' in record: raise InputError('Duplicate reserved field.')
    record['inputs']=inputs; record['input_metadata']=metas
    return record
